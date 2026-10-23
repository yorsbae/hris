"""Putaran 58 (diubah oleh patch Master Tanggungan) — pembagian beban Tagihan Mitra menurut PENANGGUNG: BPJS → perusahaan 75% / karyawan 25%; Non BPJS → 100% perusahaan (bawaan; persen bisa diubah
per tanggal lewat Master Tanggungan, lihat test_coverage). Mencakup aturan hitung, snapshot saat dicatat, form, rekap, ekspor, impor, dan cetak PDF (Daftar Karyawan Berobat + REKAP)."""
import io
from datetime import date
from decimal import Decimal as D

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.urls import reverse
from openpyxl import load_workbook

from apps.hrd import services as hrd_services
from apps.poli import services
from apps.poli import test_billing as tb
from apps.poli.models import PartnerBill


class SplitRuleTests(tb.Base):
    def test_bpjs_company_75_employee_25(self):
        self.assertEqual(services.split_bill(D(100000), "bpjs", tb.SVC), (D(75000), D(25000)))

    def test_non_bpjs_company_pays_everything(self):
        self.assertEqual(services.split_bill(D(100000), "non_bpjs", tb.SVC), (D(100000), D(0)))

    def test_parts_always_add_up_to_net_even_with_rounding(self):
        for net in (D(1), D(3), D(100001), D(181500), D("167887.50"), D(999999)):
            c, e = services.split_bill(net, "bpjs", tb.SVC); self.assertEqual(c + e, net, net); self.assertGreaterEqual(e, 0)

    @override_settings(POLI_BPJS_COMPANY_PCT=80)
    def test_default_bpjs_percentage_is_configurable(self):
        self.assertEqual(services.split_bill(D(100000), "bpjs", tb.SVC), (D(80000), D(20000)))
        self.assertEqual(services.split_bill(D(100000), "non_bpjs", tb.SVC), (D(100000), D(0)))        # Non BPJS tidak ikut pengaturan ini

    def test_has_bpjs_uses_hrd_status_first_then_number(self):
        self.assertTrue(services.employee_has_bpjs(self.e1))                      # tanpa catatan status: nomor BPJS Kes terisi
        self.assertFalse(services.employee_has_bpjs(self.e2))                     # tanpa nomor
        hrd_services.set_bpjs_status(self.e1, "kes", "nonaktif", date.today(), "keluar BPJS", self.hrd)
        self.assertFalse(services.employee_has_bpjs(self.e1))                     # status HRD nonaktif menang atas nomor
        hrd_services.set_bpjs_status(self.e2, "kes", "aktif", date.today(), "daftar", self.hrd)
        self.assertTrue(services.employee_has_bpjs(self.e2))                      # status HRD aktif walau nomor kosong


class BillCreationTests(tb.Base):
    def test_bpjs_bill_snapshots_split_pct_and_discount(self):
        b = services.create_partner_bill(self.poli, self.partner, "INV-S1", tb.BILLD, self.e1, tb.SVC, "rawat_jalan", "", None, "bpjs", D(181500), (), D(13613))
        self.assertTrue(b.has_bpjs); self.assertEqual(b.net_amount, D(167887)); self.assertEqual(b.company_amount + b.employee_amount, b.net_amount); self.assertEqual(b.company_pct, D(75))
        self.assertEqual((b.company_amount, b.employee_amount), (D(125915), D(41972)))

    def test_non_bpjs_bill_is_fully_company(self):
        b = services.create_partner_bill(self.poli, self.partner, "INV-S2", tb.BILLD, self.e2, tb.SVC, "rawat_jalan", "", None, "non_bpjs", D(131200), (), D(9840))
        self.assertFalse(b.has_bpjs); self.assertEqual((b.company_amount, b.employee_amount, b.company_pct), (D(121360), D(0), D(100)))

    def test_payer_decides_the_split_not_the_employee_bpjs_status(self):
        b = services.create_partner_bill(self.poli, self.partner, "INV-S5", tb.BILLD, self.e2, tb.SVC, "rawat_jalan", "", None, "bpjs", D(100000), ())   # e2 tanpa BPJS, tetapi penanggung BPJS
        self.assertEqual((b.has_bpjs, b.company_amount, b.employee_amount), (False, D(75000), D(25000)))
        b = services.create_partner_bill(self.poli, self.partner, "INV-S6", tb.BILLD, self.e1, tb.SVC, "rawat_jalan", "", None, "non_bpjs", D(100000), ())  # e1 ber-BPJS, tetapi penanggung Non BPJS
        self.assertEqual((b.has_bpjs, b.company_amount, b.employee_amount), (True, D(100000), D(0)))

    def test_snapshot_does_not_move_when_bpjs_changes_later(self):
        b = services.create_partner_bill(self.poli, self.partner, "INV-S3", tb.BILLD, self.e1, tb.SVC, "rawat_jalan", "", None, "bpjs", D(100000), ())
        hrd_services.set_bpjs_status(self.e1, "kes", "nonaktif", date.today(), "keluar", self.hrd)
        b.refresh_from_db(); self.assertEqual((b.has_bpjs, b.company_amount, b.employee_amount), (True, D(75000), D(25000)))

    def test_old_payer_values_are_rejected(self):
        for old in ("perusahaan", "karyawan"):
            r = self.post(bill_number="INV-OLD-" + old, payer=old); self.assertEqual(r.status_code, 200); self.assertIn("payer", r.context["form"].errors)

    def test_discount_cannot_exceed_total(self):
        with self.assertRaisesMessage(ValueError, "Potongan"):
            services.create_partner_bill(self.poli, self.partner, "INV-S4", tb.BILLD, self.e1, tb.SVC, "rawat_jalan", "", None, "bpjs", D(1000), (), D(1001))

    def test_form_creates_with_discount_and_rejects_excess(self):
        r = self.post(bill_number="INV-F1", total_amount="181.500", discount_amount="13.613"); self.assertEqual(r.status_code, 302, getattr(r, "context", None) and r.context["form"].errors)
        b = PartnerBill.objects.get(bill_number="INV-F1"); self.assertEqual((b.discount_amount, b.company_amount, b.employee_amount), (D(13613), D(125915), D(41972)))
        r = self.post(bill_number="INV-F2", total_amount="1.000", discount_amount="2.000"); self.assertEqual(r.status_code, 200); self.assertIn("discount_amount", r.context["form"].errors)

    def test_detail_shows_split(self):
        b = self.bill(employee=self.e1, total=D(100000))
        r = self.client.get(reverse("poli_bill_detail", args=[b.pk])); self.assertContains(r, "Ditanggung perusahaan"); self.assertContains(r, "Rp 75.000"); self.assertContains(r, "Rp 25.000")


class RekapAndExportTests(tb.Base):
    def setUp(self):
        super().setUp()
        self.b1 = services.create_partner_bill(self.poli, self.partner, "R1", tb.BILLD, self.e1, tb.SVC, "rawat_jalan", "", None, "bpjs", D(181500), (), D(13613))   # penanggung BPJS, Produksi
        self.b2 = services.create_partner_bill(self.poli, self.partner, "R2", tb.BILLD, self.e2, tb.SVC, "rawat_jalan", "", None, "non_bpjs", D(131200), (), D(9840))    # Non BPJS, Gudang
        self.bx = services.create_partner_bill(self.poli, self.partner, "R3", tb.BILLD, self.e1, tb.SVC, "rawat_jalan", "", None, "bpjs", D(900000), ())
        self.client.post(reverse("poli_bill_action", args=[self.bx.pk, "reject"]), {"note": "dobel"})

    def get(self, **q): return self.client.get(reverse("poli_bills"), {"period": tb.PER, **q})

    def test_totals_split_exclude_rejected(self):
        r = self.get(); s = r.context["split"]
        self.assertEqual((s["company"], s["employee"]), (D(125915) + D(121360), D(41972))); self.assertEqual(s["disc"], D(13613) + D(9840))
        self.assertContains(r, 'id="split-company"'); self.assertContains(r, 'id="split-employee"')
        dept = {x["employee__department__name"]: x for x in r.context["by_dept"]}; self.assertEqual((dept["Produksi"]["company"], dept["Produksi"]["employee"]), (D(125915), D(41972))); self.assertEqual(dept["Gudang"]["employee"], D(0))

    def test_export_has_split_columns_and_values(self):
        r = self.get(export="summary", format="xlsx"); ws = load_workbook(io.BytesIO(r.content)).active; rows = list(ws.iter_rows(values_only=True)); h = list(rows[0])
        for c in ("potongan", "tagihan_aktif", "bpjs_aktif", "persen_perusahaan", "ditanggung_perusahaan", "ditanggung_karyawan"): self.assertIn(c, h)
        by = {row[h.index("nomor_tagihan")]: row for row in rows[1:]}
        self.assertEqual(set(by), {"R1", "R2", "R3"}); r1 = by["R1"]
        self.assertEqual((r1[h.index("tagihan_aktif")], r1[h.index("ditanggung_perusahaan")], r1[h.index("ditanggung_karyawan")], r1[h.index("bpjs_aktif")]), (167887, 125915, 41972, "ya"))
        self.assertEqual(by["R2"][h.index("bpjs_aktif")], "tidak"); self.assertEqual((r1[h.index("persen_perusahaan")], by["R2"][h.index("persen_perusahaan")]), (75, 100))

    def test_print_pdf_is_poli_only_and_excludes_rejected(self):
        r = self.get(print="1"); self.assertEqual(r.status_code, 200); self.assertEqual(r["Content-Type"], "application/pdf"); self.assertTrue(r.content.startswith(b"%PDF"))
        from apps.core.models import AuditLog
        self.assertTrue(AuditLog.objects.filter(module="poli", action="bill_print").exists())
        for who in ("hrd", "adm"):
            self.login(who); self.assertEqual(self.get(print="1").status_code, 403)

    def test_print_pdf_content_lists_split_and_rekap(self):
        from unittest import mock
        from reportlab.pdfgen import canvas as rl_canvas
        from apps.poli import pdf
        bills = list(PartnerBill.objects.exclude(status="ditolak").select_related("employee__position", "employee__department").order_by("id"))
        texts, orig = [], rl_canvas.Canvas.drawString
        def spy(c, x, y, t, *a, **k): texts.append(str(t)); return orig(c, x, y, t, *a, **k)
        with mock.patch.object(rl_canvas.Canvas, "drawString", spy): raw = pdf.bills_pdf(bills, tb.PER, self.partner, self.poli)
        self.assertTrue(raw.startswith(b"%PDF")); self.assertGreater(len(raw), 2000)
        joined = "\n".join(texts)
        self.assertIn("BPJS: prs 75% / kry 25%", joined); self.assertIn("Non BPJS: prs 100% / kry 0%", joined)   # keterangan per baris mengikuti penanggung
        self.assertIn("Dicetak", joined); self.assertNotIn("Beban", joined)                                        # kalimat "Beban: karyawan ber-BPJS …" sudah dihapus dari cetakan


class ImportTests(tb.Base):
    def test_import_with_potongan_column_applies_split(self):
        hdr = "mitra,nomor_tagihan,tanggal_tagihan,nik,tanggal_layanan,jenis_layanan,kode_diagnosa,penanggung,total,potongan,keluhan\n"
        row = f"RS Sehat,IMP-1,{tb.BILLD.isoformat()},001,{tb.SVC.isoformat()},rawat_jalan,A09,bpjs,181500,13613,\n"
        r = self.client.post(reverse("poli_bill_import"), {"file": SimpleUploadedFile("t.csv", (hdr + row).encode()), "mode": "import"}); self.assertEqual(r.status_code, 200)
        b = PartnerBill.objects.get(bill_number="IMP-1"); self.assertEqual((b.discount_amount, b.company_amount, b.employee_amount), (D(13613), D(125915), D(41972)))

    def test_reimport_same_file_does_not_duplicate_bills(self):
        hdr = "mitra,nomor_tagihan,tanggal_tagihan,nik,tanggal_layanan,jenis_layanan,kode_diagnosa,penanggung,total,potongan,keluhan\n"
        row = f"RS Sehat,IMP-2,{tb.BILLD.isoformat()},002,{tb.SVC.isoformat()},rawat_jalan,A09,bpjs,100000,0,\n"
        up = lambda: self.client.post(reverse("poli_bill_import"), {"file": SimpleUploadedFile("t.csv", (hdr + row).encode()), "mode": "import"})  # noqa: E731
        up(); r = up()
        self.assertEqual(PartnerBill.objects.filter(bill_number="IMP-2").count(), 1); self.assertContains(r, "sudah tercatat")
