"""Tagihan Mitra (putaran 24, P6): RBAC data medis, validasi, alur status, enkripsi keluhan, audit tanpa isi medis, rekap, ekspor, impor, mitra."""
import io
from datetime import date, timedelta
from decimal import Decimal
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.urls import reverse
from openpyxl import load_workbook
from apps.core.models import AuditLog
from apps.core.money import XLSX_RUPIAH_FORMAT
from apps.poli import services
from apps.poli.models import Diagnosis, Partner, PartnerBill, PartnerBillEvent
from apps.poli.test_poli import PoliBase

TODAY = date.today()
SVC = TODAY - timedelta(days=3)
BILLD = TODAY - timedelta(days=1)
PER = SVC.strftime("%Y-%m")
SECRET = "nyeri dada hebat RAHASIA-MEDIS"
HDR = "mitra,nomor_tagihan,tanggal_tagihan,nik,tanggal_layanan,jenis_layanan,kode_diagnosa,penanggung,total,keluhan\n"


class Base(PoliBase):
    def setUp(self):
        self.login("poli"); self.partner = Partner.objects.create(name="RS Sehat", kind="rs"); self.partner2 = Partner.objects.create(name="Klinik Maju", kind="klinik")

    def lines(self, *rows):
        d = {"ln-TOTAL_FORMS": "6", "ln-INITIAL_FORMS": "0", "ln-MIN_NUM_FORMS": "0", "ln-MAX_NUM_FORMS": "40"}
        for i in range(6):
            desc, q, p = rows[i] if i < len(rows) else ("", "", "")
            d.update({f"ln-{i}-description": desc, f"ln-{i}-qty": q, f"ln-{i}-unit_price": p})
        return d

    def post(self, lines=(), **kw):
        d = {"partner": self.partner.pk, "bill_number": "INV-1", "bill_date": BILLD.isoformat(), "nik": "001", "service_date": SVC.isoformat(), "service_type": "rawat_jalan",
             "complaint": SECRET, "diagnosis_code": "A09", "payer": "perusahaan", "total_amount": "450.000", **self.lines(*lines)}; d.update(kw)
        return self.client.post(reverse("poli_bill_new"), d)

    def bill(self, **kw):
        d = dict(user=self.poli, partner=self.partner, bill_number="B-%d" % (PartnerBill.objects.count() + 1), bill_date=BILLD, employee=self.e1, service_date=SVC, service_type="rawat_jalan",
                 complaint=SECRET, diagnosis=self.diag, payer="perusahaan", total=Decimal(450000)); d.update(kw)
        return services.create_partner_bill(d["user"], d["partner"], d["bill_number"], d["bill_date"], d["employee"], d["service_date"], d["service_type"], d["complaint"], d["diagnosis"], d["payer"], d["total"], d.get("lines", ()))

    def act(self, b, action, **kw): return self.client.post(reverse("poli_bill_action", args=[b.pk, action]), kw)


class AccessTests(Base):
    def test_poli_and_superadmin_ok_others_forbidden(self):
        b = self.bill()
        urls = [reverse(n) for n in ("poli_bills", "poli_bill_new", "poli_partners", "poli_bill_import")] + [reverse("poli_bill_detail", args=[b.pk]), reverse("poli_bills") + "?export=detail"]
        for who, code in (("poli", 200), ("su", 200), ("hrd", 403), ("adm", 403)):
            self.login(who)
            for u in urls: self.assertEqual(self.client.get(u).status_code, code, (who, u))

    def test_forbidden_roles_cannot_write(self):
        b = self.bill()
        for who in ("hrd", "adm"):
            self.login(who)
            self.assertEqual(self.post(bill_number="X").status_code, 403)
            for a in ("verify", "approve", "reject", "pay"): self.assertEqual(self.act(b, a, note="x").status_code, 403, (who, a))
            self.assertEqual(self.client.post(reverse("poli_partners"), {"action": "add", "name": "Z", "kind": "rs"}).status_code, 403)
        b.refresh_from_db(); self.assertEqual(b.status, "diterima"); self.assertEqual(PartnerBill.objects.count(), 1); self.assertFalse(Partner.objects.filter(name="Z").exists())

    def test_anonymous_redirected_and_actions_post_only(self):
        b = self.bill(); self.client.logout()
        for n in ("poli_bills", "poli_bill_new", "poli_partners"): self.assertEqual(self.client.get(reverse(n)).status_code, 302)
        self.login("poli"); self.assertEqual(self.client.get(reverse("poli_bill_action", args=[b.pk, "verify"])).status_code, 405)

    def test_hub_and_sidebar_for_poli_only(self):
        self.assertContains(self.client.get("/poli/"), "/poli/billing/")
        items = [i["label"] for g in self.client.get("/poli/").context["nav_groups"] for i in g["items"]]; self.assertIn("Tagihan Mitra", items)
        self.login("hrd"); items = [i["label"] for g in self.client.get("/").context["nav_groups"] for i in g["items"]]; self.assertNotIn("Tagihan Mitra", items)


class CreateTests(Base):
    def test_create_with_lines_and_event_and_clean_audit(self):
        r = self.post(lines=[("Konsultasi", "1", "150.000"), ("Obat", "2", "150.000")]); self.assertEqual(r.status_code, 302)
        b = PartnerBill.objects.get(); self.assertEqual((b.total_amount, b.lines_total, b.status, b.diagnosis.code, b.created_by.username), (Decimal(450000), Decimal(450000), "diterima", "A09", "poli"))
        self.assertFalse(b.mismatch); self.assertEqual([l.amount for l in b.lines.all()], [Decimal(150000), Decimal(300000)])
        self.assertEqual([(e.from_status, e.to_status) for e in b.events.all()], [("", "diterima")])
        blob = str(list(AuditLog.objects.filter(module="poli").values_list("before", "after")))
        self.assertNotIn("RAHASIA", blob); self.assertNotIn("A09", blob); self.assertNotIn("Diare", blob)
        a = AuditLog.objects.filter(action="bill_create").get(); self.assertEqual((a.after["number"], a.after["employee"], a.after["total"]), ("INV-1", "001", "450000"))

    def test_complaint_is_encrypted_at_rest_but_readable(self):
        self.post(); b = PartnerBill.objects.get()
        with connection.cursor() as c: c.execute("SELECT complaint FROM poli_partnerbill WHERE id = %s", [b.pk]); raw = c.fetchone()[0]
        self.assertNotIn("RAHASIA", raw); self.assertEqual(PartnerBill.objects.get(pk=b.pk).complaint, SECRET)
        self.assertContains(self.client.get(reverse("poli_bill_detail", args=[b.pk])), "RAHASIA-MEDIS")

    def test_mismatch_flagged_not_rejected(self):
        self.post(lines=[("Konsultasi", "1", "100.000")]); b = PartnerBill.objects.get(); self.assertTrue(b.mismatch)
        self.assertContains(self.client.get(reverse("poli_bills"), {"period": PER}), "selisih")

    def test_no_lines_is_fine(self):
        self.post(); b = PartnerBill.objects.get(); self.assertEqual((b.lines.count(), b.lines_total, b.mismatch), (0, 0, False))

    def test_partial_line_rejected(self):
        self.assertEqual(self.post(lines=[("Konsultasi", "1", "")]).status_code, 200); self.assertEqual(self.post(lines=[("", "1", "5000")]).status_code, 200)
        self.assertEqual(PartnerBill.objects.count(), 0)

    def test_invalid_values_rejected(self):
        for bad in ({"total_amount": "0"}, {"total_amount": "-1"}, {"total_amount": "1.5"}, {"nik": "999"}, {"diagnosis_code": "ZZZ"}, {"service_type": "x"}, {"payer": "x"}, {"partner": "999"},
                    {"service_date": (TODAY + timedelta(days=1)).isoformat()}, {"service_date": (TODAY + timedelta(days=1)).isoformat(), "bill_date": (TODAY + timedelta(days=2)).isoformat()}, {"bill_date": (SVC - timedelta(days=1)).isoformat()}, {"bill_number": " "}, {"service_date": "bukan"}):
            self.assertEqual(self.post(**bad).status_code, 200, bad)
        self.assertEqual(PartnerBill.objects.count(), 0)

    def test_inactive_employee_allowed_and_diagnosis_optional(self):
        self.assertEqual(self.post(nik="009", diagnosis_code="").status_code, 302); self.assertIsNone(PartnerBill.objects.get().diagnosis)

    def test_inactive_partner_not_selectable(self):
        self.partner.is_active = False; self.partner.save(); self.assertEqual(self.post().status_code, 200); self.assertEqual(PartnerBill.objects.count(), 0)

    def test_duplicate_number_same_partner_rejected_case_space_insensitive_but_other_partner_ok(self):
        self.post(); self.assertEqual(self.post(bill_number="  inv-1 ").status_code, 200); self.assertEqual(PartnerBill.objects.count(), 1)
        self.assertEqual(self.post(partner=self.partner2.pk).status_code, 302); self.assertEqual(PartnerBill.objects.count(), 2)

    def test_rejected_number_can_be_reused(self):
        self.post(); b = PartnerBill.objects.get(); self.act(b, "reject", note="salah input")
        self.assertEqual(self.post().status_code, 302); self.assertEqual(PartnerBill.objects.count(), 2)

    def test_immutable_rows(self):
        b = self.bill(lines=[("x", 1, 450000)])
        with self.assertRaises(PermissionError): b.delete()
        ev = b.events.first(); ev.note = "ubah"
        with self.assertRaises(PermissionError): ev.save()
        with self.assertRaises(PermissionError): ev.delete()


class FlowTests(Base):
    def test_happy_path_with_events_and_audit(self):
        b = self.bill()
        self.act(b, "verify"); self.act(b, "approve")
        r = self.act(b, "pay", paid_date=TODAY.isoformat(), payment_ref="TRF-77"); self.assertEqual(r.status_code, 302)
        b.refresh_from_db(); self.assertEqual((b.status, b.paid_date, b.payment_ref), ("dibayar", TODAY, "TRF-77"))
        self.assertEqual([e.to_status for e in b.events.all()], ["diterima", "diverifikasi", "disetujui", "dibayar"]); self.assertTrue(all(e.user.username == "poli" for e in b.events.all()))
        for a in ("bill_verify", "bill_approve", "bill_pay"): self.assertTrue(AuditLog.objects.filter(action=a).exists(), a)
        self.assertNotIn("RAHASIA", str(list(AuditLog.objects.values_list("before", "after"))))

    def test_illegal_transitions_blocked(self):
        b = self.bill()
        for a, kw in (("approve", {}), ("pay", {"paid_date": TODAY.isoformat(), "payment_ref": "x"})):
            self.act(b, a, **kw); b.refresh_from_db(); self.assertEqual(b.status, "diterima", a)
        self.act(b, "verify"); self.act(b, "verify"); b.refresh_from_db(); self.assertEqual(b.status, "diverifikasi"); self.assertEqual(b.events.count(), 2)       # verify ulang tidak menambah jejak
        self.act(b, "pay", paid_date=TODAY.isoformat(), payment_ref="x"); b.refresh_from_db(); self.assertEqual(b.status, "diverifikasi")

    def test_terminal_states_are_final(self):
        b = self.bill(); self.act(b, "reject", note="dobel")
        for a in ("verify", "approve", "reject", "pay"): self.act(b, a, note="x", paid_date=TODAY.isoformat(), payment_ref="x")
        b.refresh_from_db(); self.assertEqual(b.status, "ditolak")
        c = self.bill(); self.act(c, "verify"); self.act(c, "approve"); self.act(c, "pay", paid_date=TODAY.isoformat(), payment_ref="x"); self.act(c, "reject", note="x")
        c.refresh_from_db(); self.assertEqual(c.status, "dibayar")

    def test_reject_needs_reason_from_every_open_state(self):
        for steps in ((), ("verify",), ("verify", "approve")):
            b = self.bill()
            for s in steps: self.act(b, s)
            self.act(b, "reject", note="  "); b.refresh_from_db(); self.assertNotEqual(b.status, "ditolak", steps)
            self.act(b, "reject", note="tagihan dobel"); b.refresh_from_db(); self.assertEqual((b.status, b.status_note), ("ditolak", "tagihan dobel"), steps)

    def test_pay_validation(self):
        b = self.bill(); self.act(b, "verify"); self.act(b, "approve")
        for kw in ({"paid_date": TODAY.isoformat(), "payment_ref": " "}, {"payment_ref": "x"}, {"paid_date": "bukan", "payment_ref": "x"},
                   {"paid_date": (TODAY + timedelta(days=1)).isoformat(), "payment_ref": "x"}, {"paid_date": (BILLD - timedelta(days=1)).isoformat(), "payment_ref": "x"}):
            self.act(b, "pay", **kw); b.refresh_from_db(); self.assertEqual(b.status, "disetujui", kw)

    def test_unknown_action_and_pk(self):
        b = self.bill(); self.assertEqual(self.act(b, "hapus").status_code, 302); b.refresh_from_db(); self.assertEqual(b.status, "diterima")
        self.assertEqual(self.client.post(reverse("poli_bill_action", args=[99999, "verify"])).status_code, 404); self.assertEqual(self.client.get(reverse("poli_bill_detail", args=[99999])).status_code, 404)

    def test_detail_view_is_audited_without_content(self):
        b = self.bill(); self.client.get(reverse("poli_bill_detail", args=[b.pk]))
        a = AuditLog.objects.filter(action="bill_view").get(); self.assertEqual(a.object_id, str(b.pk)); self.assertNotIn("RAHASIA", str((a.before, a.after)))

    def test_detail_hides_actions_not_available(self):
        b = self.bill(); r = self.client.get(reverse("poli_bill_detail", args=[b.pk])); self.assertContains(r, "Tandai diverifikasi"); self.assertNotContains(r, "Catat dibayar")
        self.act(b, "reject", note="x"); r = self.client.get(reverse("poli_bill_detail", args=[b.pk])); self.assertNotContains(r, "Tandai diverifikasi"); self.assertContains(r, "Alasan ditolak")


class RecapTests(Base):
    def setUp(self):
        super().setUp()
        self.b1 = self.bill(total=Decimal(100000))                                                   # Produksi, RS Sehat, A09
        self.b2 = self.bill(employee=self.e2, total=Decimal(250000), partner=self.partner2, diagnosis=None)  # Gudang, Klinik Maju
        self.b3 = self.bill(employee=self.e3, total=Decimal(50000))                                  # Produksi, RS Sehat
        self.bx = self.bill(total=Decimal(900000)); self.act(self.bx, "reject", note="dobel")        # ditolak: tak dihitung
        self.act(self.b1, "verify"); self.act(self.b1, "approve"); self.act(self.b1, "pay", paid_date=TODAY.isoformat(), payment_ref="T1")

    def get(self, **q): return self.client.get(reverse("poli_bills"), {"period": PER, **q})

    def test_totals_exclude_rejected_and_outstanding(self):
        r = self.get(); self.assertEqual(r.context["total_live"], Decimal(400000)); self.assertEqual(r.context["outstanding"], Decimal(300000))
        st = {s["label"]: s for s in r.context["by_status"]}; self.assertEqual((st["Ditolak"]["n"], st["Ditolak"]["total"]), (1, Decimal(900000))); self.assertEqual(st["Dibayar"]["n"], 1)
        self.assertContains(r, "Rp 400.000"); self.assertNotContains(r, ",00")

    def test_breakdowns(self):
        r = self.get(); p = {x["partner__name"]: x for x in r.context["by_partner"]}
        self.assertEqual((p["RS Sehat"]["n"], p["RS Sehat"]["total"]), (2, Decimal(150000))); self.assertEqual(p["Klinik Maju"]["total"], Decimal(250000))
        d = {x["employee__department__name"]: x for x in r.context["by_dept"]}; self.assertEqual((d["Produksi"]["n"], d["Produksi"]["total"]), (2, Decimal(150000))); self.assertEqual(d["Gudang"]["total"], Decimal(250000))
        dg = {x["diagnosis__code"]: x for x in r.context["by_diag"]}; self.assertEqual(dg["A09"]["total"], Decimal(150000)); self.assertEqual(dg[None]["total"], Decimal(250000))
        self.assertEqual([x["employee__nik"] for x in r.context["employees"]], ["002", "001", "003"])

    def test_filters(self):
        self.assertEqual(self.get(partner=self.partner2.pk).context["page"].paginator.count, 1); self.assertEqual(self.get(status="dibayar").context["page"].paginator.count, 1)
        self.assertEqual(self.get(department=self.d2.pk).context["page"].paginator.count, 1); self.assertEqual(self.get(diagnosis="a09").context["page"].paginator.count, 3)
        self.assertEqual(self.get(q="INV").context["page"].paginator.count, 0); self.assertEqual(self.get(q="sari").context["page"].paginator.count, 1)
        self.assertEqual(self.get(period="2001-01").context["page"].paginator.count, 0); self.assertEqual(self.client.get(reverse("poli_bills"), {"period": ""}).context["page"].paginator.count, 4)

    def test_default_period_is_this_month_of_service_date(self):
        self.assertEqual(self.client.get(reverse("poli_bills")).context["f"]["period"], TODAY.strftime("%Y-%m"))

    def test_garbage_filters_do_not_500(self):
        for q in ("period=zzz", "partner=abc", "department=%00", "status=%27", "payer=x", "page=999", "export=zzz", "diagnosis=%00"):
            self.assertIn(self.client.get(reverse("poli_bills") + "?" + q).status_code, (200, 400), q)


class ExportTests(Base):
    def setUp(self):
        super().setUp(); self.bill(total=Decimal(100000)); self.bill(employee=self.e2, total=Decimal(250000))

    def sheet(self, kind, **q):
        r = self.client.get(reverse("poli_bills"), {"period": PER, "export": kind, "format": "xlsx", **q}); self.assertEqual(r.status_code, 200)
        return list(load_workbook(io.BytesIO(r.content)).worksheets[0].iter_rows())

    def test_summary_has_no_medical_columns_or_content(self):
        rows = self.sheet("summary"); head = [c.value for c in rows[0]]
        self.assertNotIn("keluhan", head); self.assertNotIn("diagnosa", head)
        self.assertNotIn(b"RAHASIA", self.client.get(reverse("poli_bills"), {"period": PER, "export": "summary"}).content)

    def test_detail_has_medical_columns_and_numeric_money(self):
        rows = self.sheet("detail"); head = [c.value for c in rows[0]]; self.assertEqual(head[-2:], ["diagnosa", "keluhan"])
        self.assertIn(SECRET, [r[head.index("keluhan")].value for r in rows[1:]])
        i = head.index("total"); self.assertEqual(sum(r[i].value for r in rows[1:]), 350000); self.assertEqual(rows[1][i].number_format, XLSX_RUPIAH_FORMAT)

    def test_exports_audited_without_medical_content(self):
        for k in ("summary", "detail"): self.client.get(reverse("poli_bills"), {"period": PER, "export": k})
        self.assertEqual(AuditLog.objects.filter(action="bill_export").count(), 2); self.assertNotIn("RAHASIA", str(list(AuditLog.objects.values_list("before", "after"))))

    def test_forbidden_roles_cannot_export(self):
        for who in ("hrd", "adm"):
            self.login(who)
            for k in ("summary", "detail"): self.assertEqual(self.client.get(reverse("poli_bills"), {"export": k}).status_code, 403)

    def test_formula_like_text_neutralised_in_csv(self):
        self.bill(complaint="=HYPERLINK(\"x\")"); b = self.client.get(reverse("poli_bills"), {"period": PER, "export": "detail"}).content.decode("utf-8-sig"); self.assertIn("'=HYPERLINK", b)


class ImportTests(Base):
    def upload(self, text, mode="import"): return self.client.post(reverse("poli_bill_import"), {"file": SimpleUploadedFile("t.csv", text.encode()), "mode": mode})

    def row(self, **kw):
        d = {"mitra": "RS Sehat", "no": "INV-9", "td": BILLD.isoformat(), "nik": "001", "sd": SVC.isoformat(), "jl": "rawat_jalan", "kd": "A09", "py": "perusahaan", "tot": "450000", "kel": "batuk"}; d.update(kw)
        return ",".join([d["mitra"], d["no"], d["td"], d["nik"], d["sd"], d["jl"], d["kd"], d["py"], d["tot"], d["kel"]]) + "\n"

    def test_check_only_saves_nothing(self):
        self.assertContains(self.upload(HDR + self.row(), "check"), "valid"); self.assertEqual(PartnerBill.objects.count(), 0)

    def test_import_creates_bills_with_events_and_summary_audit(self):
        self.upload(HDR + self.row() + self.row(no="INV-10", jl="Rawat Inap", py="BPJS", tot="1.200.000"))
        self.assertEqual(PartnerBill.objects.count(), 2); b = PartnerBill.objects.get(bill_number="INV-10"); self.assertEqual((b.service_type, b.payer, b.total_amount), ("rawat_inap", "bpjs", Decimal(1200000)))
        self.assertTrue(all(x.events.count() == 1 for x in PartnerBill.objects.all())); self.assertEqual(PartnerBill.objects.get(bill_number="INV-9").complaint, "batuk")
        self.assertTrue(AuditLog.objects.filter(action="tagihan-mitra_import").exists()); self.assertNotIn("batuk", str(list(AuditLog.objects.values_list("before", "after"))))

    def test_reimport_rejected_and_all_or_nothing(self):
        self.upload(HDR + self.row()); r = self.upload(HDR + self.row()); self.assertEqual(PartnerBill.objects.count(), 1); self.assertContains(r, "sudah tercatat")
        for body in (self.row(no="N1") + self.row(no="N2", nik="999"), self.row(mitra="RS Tak Dikenal"), self.row(kd="ZZZ"), self.row(tot="0"), self.row(kel="=1+1"),
                     self.row(no="D1") + self.row(no="d1"), self.row(sd=(TODAY + timedelta(days=2)).isoformat())):
            before = PartnerBill.objects.count(); self.upload(HDR + body); self.assertEqual(PartnerBill.objects.count(), before, body)

    def test_template_download(self):
        r = self.client.get("/poli/billing/import/template.csv"); self.assertEqual(r.status_code, 200); self.assertIn(b"nomor_tagihan", r.content)

    def test_forbidden_roles_cannot_import(self):
        for who in ("hrd", "adm"):
            self.login(who); self.assertEqual(self.upload(HDR + self.row(no="Z")).status_code, 403)
        self.assertEqual(PartnerBill.objects.count(), 0)


class PartnerTests(Base):
    def test_add_duplicate_and_toggle(self):
        post = lambda **k: self.client.post(reverse("poli_partners"), k)  # noqa: E731
        post(action="add", name="Lab Prima", kind="lab", contact="021-1"); post(action="add", name="lab  prima", kind="lab"); self.assertEqual(Partner.objects.filter(name__iexact="lab prima").count(), 1)
        p = Partner.objects.get(name="Lab Prima"); post(action="toggle", id=p.pk); p.refresh_from_db(); self.assertFalse(p.is_active); post(action="toggle", id=p.pk); p.refresh_from_db(); self.assertTrue(p.is_active)
        self.assertEqual(post(action="toggle", id=99999).status_code, 200); self.assertTrue(AuditLog.objects.filter(action="partner_add").exists())

    def test_deactivating_keeps_old_bills(self):
        b = self.bill(); self.client.post(reverse("poli_partners"), {"action": "toggle", "id": self.partner.pk}); b.refresh_from_db(); self.assertEqual(b.partner_id, self.partner.pk)
