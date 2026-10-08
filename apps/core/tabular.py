"""Impor & ekspor tabel CSV/XLSX yang dipakai bersama semua modul (putaran 18).

Ekspor : `export_response(nama, header, baris, request)` → CSV (UTF-8 BOM, terbuka benar di Excel) atau XLSX menurut `?format=csv|xlsx` (bawaan csv).
         Sel yang diawali = + - @ diberi apostrof di depan agar tidak dieksekusi sebagai rumus spreadsheet (CSV injection).
Impor  : `read_table(bytes, nama_file, wajib)` → list dict berkunci header huruf kecil. Menerima .csv (pemisah koma/titik-koma, UTF-8) dan .xlsx.
         Batas ukuran/baris; sel rumus XLSX dibaca sebagai teks nilainya (data_only) dan sel yang diawali = atau @ ditolak oleh pemanggil.
"""
import csv, io
from datetime import date, datetime, time
from django.http import HttpResponse

MAX_BYTES, MAX_ROWS = 5 * 1024 * 1024, 5000
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def fmt(request):
    return "xlsx" if request.GET.get("format") == "xlsx" else "csv"


def _cell(v):
    if v is None: return ""
    if isinstance(v, bool): return "ya" if v else "tidak"
    if isinstance(v, datetime): return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, date): return v.isoformat()
    if isinstance(v, time): return v.strftime("%H:%M")
    v = str(v)
    return "'" + v if v[:1] in ("=", "+", "-", "@") and not _is_number(v) else v


def _is_number(v):
    try: float(v); return True
    except ValueError: return False


def to_bytes(header, rows, kind="csv", sheet="Data"):
    rows = [[_cell(c) for c in r] for r in rows]
    if kind == "xlsx":
        from openpyxl import Workbook
        wb = Workbook(write_only=True); ws = wb.create_sheet(sheet[:31]); ws.append(list(header))
        for r in rows: ws.append(r)
        out = io.BytesIO(); wb.save(out); return out.getvalue(), XLSX
    out = io.StringIO(); w = csv.writer(out); w.writerow(header); w.writerows(rows)
    return ("\ufeff" + out.getvalue()).encode("utf-8"), "text/csv; charset=utf-8-sig"


def export_response(name, header, rows, request, sheet="Data"):
    kind = fmt(request); body, ctype = to_bytes(header, rows, kind, sheet)
    r = HttpResponse(body, content_type=ctype); r["Content-Disposition"] = f'attachment; filename="{name}.{kind}"'; return r


def read_table(raw, filename, required=()):
    """Return list[dict]. ValueError berisi pesan untuk pengguna."""
    if len(raw) > MAX_BYTES: raise ValueError("File terlalu besar (maks 5 MB).")
    name = (filename or "").lower()
    if name.endswith(".xlsx") or raw[:2] == b"PK": grid = _xlsx(raw)
    elif name.endswith((".xls", ".xlsm")): raise ValueError("Format .xls lama tidak didukung; simpan sebagai .xlsx atau .csv.")
    else: grid = _csv(raw)
    if not grid: raise ValueError("File tidak berisi data.")
    header = [str(h or "").strip().lower() for h in grid[0]]
    missing = [c for c in required if c not in header]
    if missing: raise ValueError("Kolom wajib tidak ada: " + ", ".join(missing))
    rows = [{h: str(v if v is not None else "").strip() for h, v in zip(header, r) if h} for r in grid[1:] if any(str(c or "").strip() for c in r)]
    if not rows: raise ValueError("File tidak berisi data.")
    if len(rows) > MAX_ROWS: raise ValueError(f"Maksimal {MAX_ROWS} baris per impor.")
    return rows


def _csv(raw):
    try: text = raw.decode("utf-8-sig")  # utf-8-sig membuang BOM dari Excel
    except UnicodeDecodeError: raise ValueError("File harus berenkode UTF-8 (di Excel: Save As → CSV UTF-8) atau gunakan .xlsx.")
    if not text.strip(): return []
    first = text.splitlines()[0]
    return list(csv.reader(io.StringIO(text), delimiter=";" if first.count(";") > first.count(",") else ","))  # Excel Indonesia memakai titik-koma


def _xlsx(raw):
    from openpyxl import load_workbook
    try: wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception: raise ValueError("File .xlsx tidak dapat dibaca (rusak atau bukan XLSX).")
    ws = wb.worksheets[0]
    def norm(v):
        if isinstance(v, datetime): return v.date().isoformat() if v.time() == time(0) else v.strftime("%Y-%m-%d %H:%M")
        if isinstance(v, date): return v.isoformat()
        if isinstance(v, time): return v.strftime("%H:%M")
        if isinstance(v, float) and v.is_integer(): return str(int(v))
        return v
    return [[norm(c) for c in r] for r in ws.iter_rows(values_only=True, max_row=MAX_ROWS + 2)]


def formula_like(v): return v[:1] in ("=", "@")  # pemanggil menolak sel seperti ini (selaras aturan impor karyawan)
