"""Putaran 25: stok seragam terintegrasi dengan pembelian (keluar), pembatalan (kembali), barang masuk, koreksi, rekap, akses, append-only."""
import io
from django.core.exceptions import PermissionDenied  # noqa: F401
from django.urls import reverse
from openpyxl import load_workbook
from apps.core.models import AuditLog, Notification
from . import services
from .models import UniformStock, UniformStockMovement
from .test_uniform import Base, PER, D0


class StockTests(Base):
    def stock(self, size=None): return UniformStock.objects.get(utype=self.type, size=size or self.L).balance

    def put_in(self, qty=10, size=None, **kw):
        return self.client.post(reverse("hrd_uniform_stock_in"), {"utype": self.type.pk, "size": (size or self.L).pk, "quantity": qty, "movement_date": D0.isoformat(), "note": "Vendor A SJ-1", **kw})

    def test_purchase_creates_keluar_movement_and_decrements(self):
        self.put_in(10); p = self.purchase(self.e1, quantity=3)
        self.assertEqual(self.stock(), 7); m = UniformStockMovement.objects.get(purchase=p, kind="keluar")
        self.assertEqual((m.quantity, m.balance_after, m.movement_date), (-3, 7, p.purchase_date))

    def test_void_returns_stock_once(self):
        self.put_in(10); p = self.purchase(self.e1, quantity=3); services.void_uniform_purchase(p.pk, "salah", self.hrd)
        self.assertEqual(self.stock(), 10); self.assertEqual(UniformStockMovement.objects.filter(purchase=p, kind="batal").count(), 1)
        with self.assertRaises(ValueError): services.void_uniform_purchase(p.pk, "lagi", self.hrd)
        self.assertEqual(self.stock(), 10)

    def test_negative_stock_allowed_and_flagged(self):
        self.purchase(self.e1, quantity=2); self.assertEqual(self.stock(), -2)
        r = self.client.get(reverse("hrd_uniform_stock")); self.assertContains(r, "stok minus"); self.assertTrue(r.context["minus"])
        row = [x for x in self.client.get(reverse("hrd_uniforms"), {"period": PER}).context["stock"] if x["size"] == "L"][0]; self.assertTrue(row["minus"])

    def test_stock_in_validation_and_audit(self):
        r = self.put_in(0); self.assertEqual(r.status_code, 200); r = self.put_in(-5); self.assertEqual(r.status_code, 200)
        self.assertFalse(UniformStockMovement.objects.exists())
        self.assertEqual(self.put_in(10).status_code, 302); self.assertEqual(self.stock(), 10); self.assertTrue(AuditLog.objects.filter(action="uniform_stock_masuk").exists())

    def test_adjust_requires_reason_and_allows_negative_delta(self):
        self.put_in(10); url = reverse("hrd_uniform_stock_adjust"); d = {"utype": self.type.pk, "size": self.L.pk, "quantity": -4, "movement_date": D0.isoformat()}
        self.assertEqual(self.client.post(url, d).status_code, 200); self.assertEqual(self.stock(), 10)       # tanpa alasan → ditolak
        self.assertEqual(self.client.post(url, {**d, "note": "rusak"}).status_code, 302); self.assertEqual(self.stock(), 6)
        self.assertEqual(self.client.post(url, {**d, "quantity": 0, "note": "x"}).status_code, 200)

    def test_card_is_append_only(self):
        self.put_in(5); m = UniformStockMovement.objects.get()
        m.note = "x"
        with self.assertRaises(PermissionError): m.save()
        with self.assertRaises(PermissionError): m.delete()

    def test_service_rejects_wrong_sign(self):
        with self.assertRaises(ValueError): services.stock_move(self.type.pk, self.L.pk, "masuk", -1, D0, self.hrd)
        with self.assertRaises(ValueError): services.stock_move(self.type.pk, self.L.pk, "keluar", 1, D0, self.hrd)

    def test_recap_columns_balance(self):
        self.put_in(10); self.purchase(self.e1, quantity=2); p2 = self.purchase(self.e2, quantity=1, size=self.M); services.void_uniform_purchase(p2.pk, "salah", self.hrd)
        rows = {r["size"]: r for r in self.client.get(reverse("hrd_uniforms"), {"period": PER}).context["stock"]}
        l = rows["L"]; self.assertEqual((l["awal"], l["masuk"], l["keluar"], l["adj"], l["saldo"]), (0, 10, 2, 0, 8))
        self.assertNotIn("M", rows)   # beli lalu batal di periode sama → netral, baris nol tidak ditampilkan

    def test_recap_start_balance_uses_card_before_period(self):
        self.put_in(10); other = "2000-01"
        rows = {r["size"]: r for r in self.client.get(reverse("hrd_uniforms"), {"period": "2099-12"}).context["stock"]}
        self.assertEqual((rows["L"]["awal"], rows["L"]["masuk"], rows["L"]["saldo"]), (10, 0, 10)); self.assertNotIn("L", {r["size"] for r in self.client.get(reverse("hrd_uniforms"), {"period": other}).context["stock"]})

    def test_form_and_csv_import_create_stock_movements(self):
        self.put_in(10); self.assertEqual(self.buy(quantity="2").status_code, 302); self.assertEqual(self.stock(), 8)
        self.upload("nik,tanggal,jenis,ukuran,jumlah,catatan\n002,%s,Seragam Kerja,L,3,\n" % D0.isoformat()); self.assertEqual(self.stock(), 5)
        self.assertEqual(UniformStockMovement.objects.filter(kind="keluar").count(), 2)

    def test_stock_export_and_card_page(self):
        self.put_in(10); self.purchase(self.e1, quantity=2)
        r = self.client.get(reverse("hrd_uniform_stock"), {"period": PER, "export": 1, "format": "xlsx"}); rows = list(load_workbook(io.BytesIO(r.content)).worksheets[0].iter_rows(values_only=True))
        self.assertEqual(rows[0][:4], ("tanggal", "jenis", "ukuran", "gerak")); self.assertEqual(len(rows) - 1, 2)
        self.assertEqual(self.client.get(reverse("hrd_uniform_stock"), {"period": PER}).status_code, 200)

    def test_only_hrd_and_superadmin(self):
        for u in ("adm", "poli"):
            self.login(u)
            for n in ("hrd_uniform_stock", "hrd_uniform_stock_in", "hrd_uniform_stock_adjust"): self.assertEqual(self.client.get(reverse(n)).status_code, 403, (u, n))
            self.assertEqual(self.client.post(reverse("hrd_uniform_stock_in"), {}).status_code, 403)


class MinStockTests(Base):
    def setmin(self, n, size=None): return self.client.post(reverse("hrd_uniform_stock"), {"utype": self.type.pk, "size": (size or self.L).pk, "min_stock": n})

    def test_set_min_audited_and_flags_low(self):
        self.client.post(reverse("hrd_uniform_stock_in"), {"utype": self.type.pk, "size": self.L.pk, "quantity": 5, "movement_date": D0.isoformat(), "note": "v"})
        self.assertEqual(self.setmin(8).status_code, 302); st = UniformStock.objects.get(utype=self.type, size=self.L)
        self.assertEqual(st.min_stock, 8); self.assertTrue(st.low); self.assertEqual(AuditLog.objects.filter(action="uniform_stock_min").count(), 1)
        r = self.client.get(reverse("hrd_uniform_stock")); self.assertEqual(r.context["low_n"], 1); self.assertContains(r, "menipis")
        self.assertEqual(self.client.get(reverse("hrd_uniforms"), {"period": PER}).context["low_n"], 1)
        self.assertEqual(UniformStockMovement.objects.count(), 1)       # minimum bukan gerak stok

    def test_zero_min_not_monitored_and_minus_not_low(self):
        self.purchase(self.e1, quantity=2); self.setmin(5); st = UniformStock.objects.get(utype=self.type, size=self.L)
        self.assertFalse(st.low)                                           # minus ditangani sebagai 'minus', bukan 'menipis'
        self.setmin(0); self.assertFalse(UniformStock.objects.get(pk=st.pk).low)

    def test_invalid_min_rejected_and_hrd_only(self):
        self.setmin(-1); self.setmin("abc"); self.assertFalse(UniformStock.objects.filter(min_stock__gt=0).exists())
        self.login("adm"); self.assertEqual(self.setmin(3).status_code, 403)


class LowStockNotifyTests(Base):
    def notes(self): return Notification.objects.filter(kind="stock")

    def setUp(self):
        super().setUp()
        self.client.post(reverse("hrd_uniform_stock_in"), {"utype": self.type.pk, "size": self.L.pk, "quantity": 5, "movement_date": D0.isoformat(), "note": "v"})
        services.set_min_stock(self.type.pk, self.L.pk, 4)

    def test_notifies_hrd_once_when_crossing_below_minimum(self):
        self.purchase(self.e1, quantity=1); self.assertFalse(self.notes().exists())              # 4 = minimum, belum di bawah
        self.purchase(self.e2, quantity=1); n = self.notes(); self.assertTrue(n.exists()); self.assertEqual(set(n.values_list("user__role", flat=True)), {"hrd"})
        self.assertIn("menipis", n[0].title); self.assertEqual(n[0].link, "/hrd/uniforms/stock/")
        c = n.count(); self.purchase(self.e3, quantity=1); self.assertEqual(self.notes().count(), c)      # terus turun → tidak menumpuk

    def test_again_after_recovery_and_read(self):
        self.purchase(self.e1, quantity=2); self.notes().update(is_read=True)
        services.stock_move(self.type.pk, self.L.pk, "masuk", 10, D0, self.hrd); self.assertEqual(self.notes().filter(is_read=False).count(), 0)
        self.purchase(self.e2, quantity=11); self.assertTrue(self.notes().filter(is_read=False).exists())

    def test_no_notification_without_minimum(self):
        services.set_min_stock(self.type.pk, self.L.pk, 0); self.purchase(self.e1, quantity=5); self.assertFalse(self.notes().exists())


class VendorOrderTests(Base):
    def test_order_list_is_shortfall_to_minimum(self):
        services.stock_move(self.type.pk, self.L.pk, "masuk", 3, D0, self.hrd); services.set_min_stock(self.type.pk, self.L.pk, 10)
        services.set_min_stock(self.type.pk, self.M.pk, 5); services.stock_move(self.type.pk, self.M.pk, "masuk", 9, D0, self.hrd)     # M cukup → tidak dipesan
        r = self.client.get(reverse("hrd_uniform_stock"), {"export": "order", "format": "xlsx"}); rows = list(load_workbook(io.BytesIO(r.content)).worksheets[0].iter_rows(values_only=True))
        self.assertEqual(rows[0], ("jenis", "ukuran", "stok_sekarang", "minimum", "jumlah_dipesan")); self.assertEqual(rows[1:], [("Seragam Kerja", "L", 3, 10, 7)])

    def test_minus_stock_orders_minimum_plus_deficit(self):
        self.purchase(self.e1, quantity=2); services.set_min_stock(self.type.pk, self.L.pk, 5)
        rows = self.client.get(reverse("hrd_uniform_stock"), {"export": "order"}).content.decode("utf-8-sig").splitlines(); self.assertIn("Seragam Kerja,L,-2,5,7", rows)
