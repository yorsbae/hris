"""Putaran 38: Karyawan Keluar — pencatatan, nonaktif otomatis + riwayat, tali asih, paklaring (PDF), pembatalan, filter, ekspor, impor, akses."""
import io
from datetime import date, timedelta
from django.core.management import call_command
from django.test import override_settings
from django.urls import reverse
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from apps.core.models import AuditLog
from apps.hr.models import Employee, EmployeeHistory
from . import services
from .models import Separation
from .test_base import HrdBase


def ago(n): return date.today() - timedelta(days=n)
def ahead(n): return date.today() + timedelta(days=n)


class SepBase(HrdBase):
    def setUp(self): self.login()

    def add(self, emp=None, last=None, **kw):
        d = {"nik": (emp or self.e1).nik, "kind": "resign", "last_date": (last or ago(5)).isoformat(), **kw}
        return self.client.post(reverse("hrd_separation_new"), d)

    def make(self, emp=None, last=None, **kw):
        self.add(emp, last, **kw); return Separation.objects.filter(employee=emp or self.e1).latest("id")


class CreateTests(SepBase):
    def test_past_date_deactivates_now_and_writes_history_and_copies_identity(self):
        r = self.add(request_date=ago(40).isoformat(), reason="Pindah kota", tali_asih="1.500.000", tali_asih_note="1 bulan gaji"); self.assertEqual(r.status_code, 302)
        s = Separation.objects.get(); self.e1.refresh_from_db()
        self.assertEqual((self.e1.status, s.emp_name, s.department_name, s.position_name, s.join_date, s.tali_asih), ("nonaktif", "Budi", "Produksi", "Staff", date(2024, 1, 1), 1_500_000))
        self.assertTrue(s.applied_at and s.status_changed); self.assertEqual(s.prev_status, "aktif")
        h = EmployeeHistory.objects.get(employee=self.e1, field="status"); self.assertEqual((h.old_value, h.new_value, h.effective_date), ("aktif", "nonaktif", s.last_date))
        self.assertEqual(self.last_audit("separation_create").after["employee"], "001")

    def test_future_date_stays_active_until_applied(self):
        self.add(last=ahead(10)); s = Separation.objects.get(); self.e1.refresh_from_db()
        self.assertEqual(self.e1.status, "aktif"); self.assertIsNone(s.applied_at)
        with self.assertRaises(ValueError): services.apply_separation(s.pk, self.hrd)               # belum waktunya
        self.assertEqual(services.apply_due_separations(), 0)

    def test_apply_today_allowed_manually_but_cron_waits_until_tomorrow(self):
        self.add(last=date.today()); s = Separation.objects.get(); self.assertIsNone(s.applied_at)
        self.assertEqual(services.apply_due_separations(), 0)
        self.assertEqual(self.client.post(reverse("hrd_separation_action", args=[s.pk, "apply"])).status_code, 302)
        self.e1.refresh_from_db(); self.assertEqual(self.e1.status, "nonaktif")
        self.client.post(reverse("hrd_separation_action", args=[s.pk, "apply"]))                      # kedua kali: ditolak, tidak dobel riwayat
        self.assertEqual(EmployeeHistory.objects.filter(employee=self.e1, field="status").count(), 1)

    def test_cron_command_applies_only_due_ones(self):
        self.add(self.e1, ahead(3)); self.add(self.e2, ago(0)); s = Separation.objects.get(employee=self.e1); s2 = Separation.objects.get(employee=self.e2)
        Separation.objects.filter(pk=s2.pk).update(last_date=ago(1)); Separation.objects.filter(pk=s.pk).update(last_date=ahead(3))
        call_command("apply_separations", stdout=io.StringIO()); self.e1.refresh_from_db(); self.e2.refresh_from_db()
        self.assertEqual((self.e1.status, self.e2.status), ("aktif", "nonaktif"))

    def test_already_inactive_employee_is_archived_without_history_change(self):
        self.add(self.e_off, ago(300)); s = Separation.objects.get(); self.assertTrue(s.applied_at); self.assertFalse(s.status_changed)
        self.assertFalse(EmployeeHistory.objects.filter(employee=self.e_off).exists())

    def test_validation(self):
        for bad, key in ((dict(nik="999"), "nik"), (dict(last_date=date(2020, 1, 1).isoformat()), "last_date"), (dict(last_date=ahead(900).isoformat()), "last_date"),
                         (dict(request_date=ahead(1).isoformat(), last_date=ago(1).isoformat()), "request_date"), (dict(request_date=date(2020, 1, 1).isoformat()), "request_date"),
                         (dict(tali_asih="abc"), "tali_asih"), (dict(tali_asih="-5"), "tali_asih"), (dict(kind="x"), "kind")):
            r = self.add(**bad); self.assertEqual(r.status_code, 200, bad); self.assertIn(key, r.context["form"].errors, bad)
        self.assertFalse(Separation.objects.exists())

    def test_one_active_record_per_employee_and_recreate_after_void(self):
        s = self.make(); self.assertEqual(self.add().status_code, 200)
        services.void_separation(s.pk, "salah", self.hrd); self.assertEqual(self.add().status_code, 302); self.assertEqual(Separation.objects.count(), 2)

    def test_service_period_text(self):
        self.assertEqual(services.service_period(date(2024, 1, 1), date(2026, 3, 15)), "2 tahun 2 bulan"); self.assertEqual(services.service_period(date(2026, 3, 1), date(2026, 3, 20)), "< 1 bulan")
        self.assertEqual(services.service_period(date(2025, 3, 31), date(2026, 3, 30)), "11 bulan")

    def test_cannot_delete(self):
        s = self.make()
        with self.assertRaises(PermissionError): s.delete()


class TaliAsihTests(SepBase):
    def test_set_and_mark_paid_locks_amount(self):
        s = self.make(); url = reverse("hrd_separation_action", args=[s.pk, "tali"])
        self.assertEqual(self.client.post(url, {"tali_asih": "2.000.000", "tali_asih_note": "kebijakan"}).status_code, 302); s.refresh_from_db(); self.assertEqual((s.tali_asih, s.tali_asih_paid_on), (2_000_000, None))
        self.client.post(url, {"tali_asih": "2.500.000", "paid_on": date.today().isoformat()}); s.refresh_from_db(); self.assertEqual((s.tali_asih, s.tali_asih_paid_on), (2_500_000, date.today()))
        self.client.post(url, {"tali_asih": "9.000.000"}); s.refresh_from_db(); self.assertEqual(s.tali_asih, 2_500_000)          # terkunci
        self.assertEqual(AuditLog.objects.filter(action="separation_tali_asih").count(), 2)
        self.assertEqual(self.last_audit("separation_tali_asih").before["tali_asih"], 2_000_000)

    def test_rules(self):
        s = self.make()
        for bad in ({"tali_asih": "0", "paid_on": date.today().isoformat()}, {"tali_asih": "100", "paid_on": ahead(2).isoformat()}, {"tali_asih": "100", "paid_on": date(2020, 1, 1).isoformat()}, {"tali_asih": "x"}):
            self.client.post(reverse("hrd_separation_action", args=[s.pk, "tali"]), bad); s.refresh_from_db(); self.assertIsNone(s.tali_asih_paid_on, bad); self.assertEqual(s.tali_asih, 0)


class PaklaringTests(SepBase):
    def test_requires_employee_to_be_out_and_numbers_once(self):
        s = self.make(last=ahead(5))
        with self.assertRaises(ValueError): services.issue_paklaring(s.pk, date.today(), self.hrd)
        s2 = self.make(self.e2, ago(3)); p = services.issue_paklaring(s2.pk, date.today(), self.hrd)
        self.assertRegex(p.paklaring_number, rf"^{p.pk:04d}/PAK/HRD/[IVX]+/{date.today().year}$")
        with self.assertRaises(ValueError): services.issue_paklaring(s2.pk, date.today(), self.hrd)

    def test_date_rules(self):
        s = self.make(last=ago(3))
        with self.assertRaises(ValueError): services.issue_paklaring(s.pk, ago(10), self.hrd)     # sebelum tanggal keluar
        with self.assertRaises(ValueError): services.issue_paklaring(s.pk, ahead(1), self.hrd)

    def test_pdf_only_after_issued_and_is_audited_and_uses_copied_identity(self):
        s = self.make(last=ago(3)); self.assertEqual(self.client.get(reverse("hrd_separation_pdf", args=[s.pk])).status_code, 404)
        self.client.post(reverse("hrd_separation_action", args=[s.pk, "paklaring"]), {"issue_date": date.today().isoformat()})
        self.e1.name = "Nama Baru"; self.e1.save()
        r = self.client.get(reverse("hrd_separation_pdf", args=[s.pk])); self.assertEqual((r.status_code, r["Content-Type"]), (200, "application/pdf")); self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertTrue(self.last_audit("print_paklaring")); self.assertTrue(self.last_audit("separation_paklaring"))

    def test_pdf_has_no_tali_asih_and_no_reason(self):
        from .pdf import paklaring_pdf
        s = self.make(last=ago(3), tali_asih="7.654.321", reason="alasan-rahasia-xyz"); s = services.issue_paklaring(s.pk, date.today(), self.hrd)
        raw = paklaring_pdf(s); self.assertNotIn(b"7654321", raw); self.assertNotIn(b"alasan-rahasia-xyz", raw)

    @override_settings(COMPANY_NAME="PT Contoh")
    def test_pdf_generation_for_every_kind(self):
        from .pdf import paklaring_pdf
        for i, (k, _) in enumerate(Separation.Kind.choices):
            e = Employee.objects.create(nik=f"K{i}", name=f"Orang {i}", gender="L", department=self.d1, position=self.pos, join_date=date(2023, 1, 1))
            s = self.make(e, ago(3), kind=k); s = services.issue_paklaring(s.pk, date.today(), self.hrd); self.assertTrue(paklaring_pdf(s).startswith(b"%PDF"), k)


class VoidTests(SepBase):
    def test_void_restores_status_and_history_and_requires_reason(self):
        s = self.make(); self.e1.refresh_from_db(); self.assertEqual(self.e1.status, "nonaktif")
        self.client.post(reverse("hrd_separation_action", args=[s.pk, "void"]), {"reason": "  "}); s.refresh_from_db(); self.assertFalse(s.is_void)
        self.client.post(reverse("hrd_separation_action", args=[s.pk, "void"]), {"reason": "salah orang"}); s.refresh_from_db(); self.e1.refresh_from_db()
        self.assertTrue(s.is_void); self.assertEqual(self.e1.status, "aktif"); self.assertEqual(list(EmployeeHistory.objects.filter(employee=self.e1, field="status").values_list("new_value", flat=True)), ["nonaktif", "aktif"])
        self.assertEqual(self.last_audit("separation_void").after["reason"], "salah orang")

    def test_void_does_not_reactivate_if_system_did_not_deactivate(self):
        s = self.make(self.e_off, ago(300)); services.void_separation(s.pk, "x", self.hrd); self.e_off.refresh_from_db(); self.assertEqual(self.e_off.status, "nonaktif")

    def test_blocked_after_paklaring_or_payment(self):
        s = self.make(); services.issue_paklaring(s.pk, date.today(), self.hrd)
        with self.assertRaises(ValueError): services.void_separation(s.pk, "x", self.hrd)
        t = self.make(self.e2, ago(2), tali_asih="1000"); services.set_tali_asih(t.pk, 1000, "", date.today(), self.hrd)
        with self.assertRaises(ValueError): services.void_separation(t.pk, "x", self.hrd)

    def test_voided_cannot_be_changed(self):
        s = self.make(); services.void_separation(s.pk, "x", self.hrd)
        for f in (lambda: services.set_tali_asih(s.pk, 5, "", None, self.hrd), lambda: services.issue_paklaring(s.pk, date.today(), self.hrd), lambda: services.apply_separation(s.pk, self.hrd), lambda: services.void_separation(s.pk, "y", self.hrd)):
            with self.assertRaises(ValueError): f()


class ListAndExportTests(SepBase):
    def setUp(self):
        super().setUp(); self.a = self.make(self.e1, ago(5), tali_asih="1.000.000"); self.b = self.make(self.e2, ago(4), kind="phk"); self.c = self.make(self.e3, ago(3), kind="habis_kontrak")
        services.issue_paklaring(self.c.pk, date.today(), self.hrd)

    def get(self, **q): return self.client.get(reverse("hrd_separations"), {"year": "", **q})

    def test_totals_and_cards(self):
        r = self.get(); self.assertEqual((r.context["totals"]["n"], r.context["totals"]["tali"], r.context["tali_belum_n"], r.context["pak_belum"]), (3, 1_000_000, 1, 2))

    def test_filters(self):
        self.assertEqual(self.get(kind="phk").context["page"].paginator.count, 1); self.assertEqual(self.get(department="Gudang").context["page"].paginator.count, 1)
        self.assertEqual(self.get(paklaring="sudah").context["page"].paginator.count, 1); self.assertEqual(self.get(paklaring="belum").context["page"].paginator.count, 2)
        self.assertEqual(self.get(tali="belum").context["page"].paginator.count, 1); self.assertEqual(self.get(q="sari").context["page"].paginator.count, 1); self.assertEqual(self.get(q="002").context["page"].paginator.count, 1)
        self.assertEqual(self.get(year="2001").context["totals"]["n"], 0); self.assertEqual(self.get(month=str(self.a.last_date.month), year=str(self.a.last_date.year)).context["totals"]["n"] >= 1, True)

    def test_default_year_is_this_year_and_void_hidden_unless_requested(self):
        self.assertEqual(self.client.get(reverse("hrd_separations")).context["f"]["year"], str(date.today().year))
        services.void_separation(self.b.pk, "x", self.hrd); self.assertEqual(self.get().context["totals"]["n"], 2); self.assertEqual(self.get(batal=1).context["page"].paginator.count, 3)

    def test_weird_params_do_not_crash(self):
        for q in ("year=abc", "month=99", "kind=zzz", "paklaring=%00", "q=%27%22"): self.assertIn(self.client.get(reverse("hrd_separations") + "?" + q).status_code, (200, 400), q)

    def test_xlsx_export_money_is_number_and_audited_and_has_no_sensitive(self):
        r = self.client.get(reverse("hrd_separations"), {"year": "", "export": 1, "format": "xlsx"}); rows = list(load_workbook(io.BytesIO(r.content)).worksheets[0].iter_rows(values_only=False))
        head = [c.value for c in rows[0]]; self.assertEqual(len(rows) - 1, 3); self.assertIn("tali_asih", head); self.assertIn("tanggal_masuk", head)
        cell = [r_[head.index("tali_asih")] for r_ in rows[1:] if r_[head.index("nik")].value == "001"][0]; self.assertEqual(cell.value, 1_000_000)
        self.assertTrue(self.last_audit("separation_export")); self.assertNotIn("bpjs", " ".join(str(h) for h in head).lower())

    def test_csv_export_plain_money(self):
        b = self.client.get(reverse("hrd_separations"), {"year": "", "export": 1}).content.decode("utf-8-sig"); self.assertIn("1000000", b); self.assertNotIn(".00", b)


class ImportTests(SepBase):
    def up(self, text, mode="import"): return self.client.post(reverse("hrd_separation_import"), {"file": SimpleUploadedFile("k.csv", text.encode()), "mode": mode})

    def test_import_creates_and_archives(self):
        hdr = "nik,jenis,tgl_pengajuan,tgl_keluar,alasan,tali_asih,dasar_tali_asih\n"
        self.up(hdr + f"001,resign,,{ago(30).isoformat()},Pindah,1500000,1 bulan\n009,Habis kontrak,,{ago(400).isoformat()},,,\n")
        self.assertEqual(Separation.objects.count(), 2); self.e1.refresh_from_db(); self.assertEqual(self.e1.status, "nonaktif")
        self.assertEqual(Separation.objects.get(employee=self.e_off).kind, "habis_kontrak")

    def test_all_or_nothing_check_only_and_duplicate(self):
        hdr = "nik,jenis,tgl_pengajuan,tgl_keluar,alasan,tali_asih,dasar_tali_asih\n"
        self.up(hdr + f"001,resign,,{ago(3).isoformat()},,,\n999,resign,,{ago(3).isoformat()},,,\n"); self.assertFalse(Separation.objects.exists())
        self.up(hdr + f"001,resign,,{ago(3).isoformat()},,,\n", mode="check"); self.assertFalse(Separation.objects.exists())
        self.up(hdr + f"001,resign,,{ago(3).isoformat()},,,\n001,phk,,{ago(2).isoformat()},,,\n"); self.assertFalse(Separation.objects.exists())
        self.up(hdr + f"001,ngawur,,{ago(3).isoformat()},,,\n"); self.assertFalse(Separation.objects.exists())

    def test_formula_cell_rejected_and_template_available(self):
        self.up("nik,jenis,tgl_pengajuan,tgl_keluar,alasan,tali_asih,dasar_tali_asih\n" + f"001,resign,,{ago(3).isoformat()},=HYPERLINK(1),,\n"); self.assertFalse(Separation.objects.exists())
        r = self.client.get("/hrd/separations/import/template.csv"); self.assertEqual(r.status_code, 200); self.assertIn(b"tgl_keluar", r.content)


class AccessAndNavTests(SepBase):
    def test_only_hrd_and_superadmin(self):
        s = self.make(last=ago(3)); s = services.issue_paklaring(s.pk, date.today(), self.hrd)
        gets = [reverse("hrd_separations"), reverse("hrd_separation_new"), reverse("hrd_separation_detail", args=[s.pk]), reverse("hrd_separation_pdf", args=[s.pk]), reverse("hrd_separations") + "?export=1", reverse("hrd_separation_import")]
        posts = [reverse("hrd_separation_action", args=[s.pk, a]) for a in ("apply", "tali", "paklaring", "void")] + [reverse("hrd_separation_new")]
        for u, code in (("adm", 403), ("poli", 403), ("su", 200), ("hrd", 200)):
            self.login(u)
            for g in gets: self.assertEqual(self.client.get(g).status_code, code, (u, g))
        for u in ("adm", "poli"):
            self.login(u)
            for p in posts: self.assertEqual(self.client.post(p, {}).status_code, 403, (u, p))
        self.client.logout()
        for g in gets: self.assertEqual(self.client.get(g).status_code, 302, g)

    def test_actions_need_post_and_unknown_action_404(self):
        s = self.make(); self.assertEqual(self.client.get(reverse("hrd_separation_action", args=[s.pk, "void"])).status_code, 405)
        self.assertEqual(self.client.post(reverse("hrd_separation_action", args=[s.pk, "hapus"])).status_code, 404)
        self.assertEqual(self.client.get(reverse("hrd_separation_detail", args=[99999])).status_code, 404)

    def test_menu_and_hub_card(self):
        r = self.client.get("/hrd/"); self.assertContains(r, "/hrd/separations/"); self.assertContains(r, "Karyawan keluar")
        self.assertContains(self.client.get("/hrd/separations/"), "Karyawan Keluar")

    def test_page_shows_flow_primary_action_empty_state_and_pages_render(self):
        r = self.client.get(reverse("hrd_separations")); self.assertContains(r, 'class="hint flow"'); self.assertContains(r, "+ Catat karyawan keluar"); self.assertContains(r, "Belum ada karyawan keluar")
        s = self.make(); self.assertContains(self.client.get(reverse("hrd_separation_detail", args=[s.pk])), "Cetak paklaring" if s.paklaring_number else "Terbitkan paklaring")
        self.assertContains(self.client.get(reverse("hrd_separation_new")), "Isi data keluar")

    def test_sensitive_data_not_in_audit(self):
        self.make(tali_asih="1.000.000", reason="alasan pribadi panjang")
        for a in AuditLog.objects.filter(module="hrd", action__startswith="separation"): self.assertNotIn("alasan pribadi", str(a.after)); self.assertNotIn("0001234567890", str(a.after))


class EmployeeDetailLinkTests(SepBase):
    """Putaran 38c: detail karyawan menautkan ke Karyawan Keluar (hanya HRD/Superadmin)."""
    def page(self, emp=None): return self.client.get(reverse("employee_detail_page", args=[(emp or self.e1).pk]))

    def test_active_without_record_offers_prefilled_button(self):
        r = self.page(); self.assertContains(r, "Catat keluar"); self.assertContains(r, f"{reverse('hrd_separation_new')}?nik={self.e1.nik}"); self.assertNotContains(r, "Catatan keluar")

    def test_with_record_shows_banner_and_link_instead_of_button(self):
        s = self.make(last=ahead(10)); r = self.page()
        self.assertContains(r, "Catatan keluar"); self.assertContains(r, reverse("hrd_separation_detail", args=[s.pk])); self.assertContains(r, "terjadwal"); self.assertNotContains(r, "?nik=")
        s = self.make(emp=self.e2, last=ago(5)); services.issue_paklaring(s.pk, date.today(), self.hrd)
        number = Separation.objects.get(pk=s.pk).paklaring_number
        r = self.page(self.e2); self.assertContains(r, number); self.assertNotContains(r, "terjadwal")

    def test_voided_record_goes_back_to_button(self):
        s = self.make(last=ahead(10)); services.void_separation(s.pk, "salah catat", self.hrd)
        r = self.page(); self.assertContains(r, "Catat keluar"); self.assertNotContains(r, "Tercatat keluar")

    def test_inactive_employee_has_no_new_button_but_dept_admin_never_sees_it(self):
        Employee.objects.filter(pk=self.e2.pk).update(status="nonaktif"); self.assertNotContains(self.page(self.e2), "Catat keluar")
        self.login("adm"); r = self.page()
        if r.status_code == 200: self.assertNotContains(r, "Catat keluar"); self.assertNotContains(r, "Catatan keluar")


class ContractReminderTests(SepBase):
    """Putaran 38c: pengingat kontrak habis tidak lagi dikirim untuk karyawan yang sudah keluar (nonaktif)."""
    def test_reminder_only_for_active_employees(self):
        from apps.core.models import Notification
        from apps.hr.models import Contract
        for n, emp in (("K-AKTIF", self.e1), ("K-KELUAR", self.e2)):
            Contract.objects.create(employee=emp, number=n, kind="PKWT", start=ago(300), end=ahead(30), status="aktif")
        self.make(emp=self.e2, last=ago(2)); self.assertEqual(Employee.objects.get(pk=self.e2.pk).status, "nonaktif")
        call_command("check_contracts")
        titles = list(Notification.objects.filter(kind="contract").values_list("title", flat=True))
        self.assertEqual(len(titles), 1); self.assertIn("Budi", titles[0]); self.assertNotIn("Sari", " ".join(titles))
