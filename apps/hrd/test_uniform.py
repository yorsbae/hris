"""Rekap Seragam (putaran 23, P5): RBAC, tarif berlaku-sejak (disalin), validasi, pembatalan/penandaan, rekap, impor, ekspor, master."""
import io
from datetime import date, timedelta
from decimal import Decimal
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import load_workbook
from apps.core.models import AuditLog
from apps.core.money import XLSX_RUPIAH_FORMAT
from . import services
from .models import UniformPurchase, UniformRate, UniformSize, UniformType
from .test_base import HrdBase

TODAY = date.today()
D0 = TODAY.replace(day=1)          # tanggal pembelian bawaan: awal bulan ini (selalu ≥ join_date 2024, ≤ hari ini)
PER = D0.strftime("%Y-%m")
HDR = "nik,tanggal,jenis,ukuran,jumlah,catatan\n"


class Base(HrdBase):
    def setUp(self):
        self.login(); self.type = UniformType.objects.get(name="Seragam Kerja"); self.L = UniformSize.objects.get(code="L"); self.M = UniformSize.objects.get(code="M")

    def buy(self, **kw):
        d = {"nik": "001", "purchase_date": D0.isoformat(), "utype": self.type.pk, "size": self.L.pk, "quantity": "2", "note": ""}; d.update(kw)
        return self.client.post(reverse("hrd_uniform_new"), d)

    def upload(self, text, mode="import", name="u.csv"):
        return self.client.post(reverse("hrd_uniform_import"), {"file": SimpleUploadedFile(name, text.encode()), "mode": mode})

    def purchase(self, e=None, **kw):
        e = e or self.e1; rate = services.uniform_rate_for(e.gender, D0); qty = kw.pop("quantity", 1)
        d = dict(employee=e, purchase_date=D0, utype=self.type, size=self.L, gender=e.gender, quantity=qty, rate_amount=rate.amount, deduction_amount=rate.amount * qty, created_by=self.hrd); d.update(kw)
        p = UniformPurchase.objects.create(**d)
        services.stock_move(p.utype_id, p.size_id, "keluar", -p.quantity, p.purchase_date, self.hrd, purchase=p)   # sama seperti form (putaran 25)
        return p


class SeedTests(Base):
    def test_seed_data_present(self):
        self.assertEqual(list(UniformSize.objects.values_list("code", flat=True)), ["S", "M", "L", "XL", "XXL", "3XL"])
        self.assertEqual({r.gender: r.amount for r in UniformRate.objects.all()}, {"L": Decimal(19000), "P": Decimal(17000)})


class AccessTests(Base):
    URLS = ["hrd_uniforms", "hrd_uniform_new", "hrd_uniform_master", "hrd_uniform_import"]

    def test_only_hrd_and_superadmin(self):
        p = self.purchase()
        for who, code in (("hrd", 200), ("su", 200), ("adm", 403), ("poli", 403)):
            self.login(who)
            for u in self.URLS: self.assertEqual(self.client.get(reverse(u)).status_code, code, (who, u))
            self.assertEqual(self.client.get(reverse("hrd_uniform_detail", args=[p.pk])).status_code, code, who)
            self.assertEqual(self.client.get(reverse("hrd_uniforms") + "?export=detail").status_code, code, who)

    def test_denied_roles_cannot_write(self):
        p = self.purchase()
        for who in ("adm", "poli"):
            self.login(who)
            self.assertEqual(self.buy().status_code, 403)
            self.assertEqual(self.client.post(reverse("hrd_uniform_void", args=[p.pk]), {"reason": "x"}).status_code, 403)
            self.assertEqual(self.client.post(reverse("hrd_uniform_mark", args=[p.pk]), {"period": PER}).status_code, 403)
            self.assertEqual(self.client.post(reverse("hrd_uniform_master"), {"action": "add_type", "name": "Z"}).status_code, 403)
            self.assertEqual(self.upload(HDR + f"001,{D0},Seragam Kerja,L,1,\n").status_code, 403)
        p.refresh_from_db(); self.assertFalse(p.is_void); self.assertEqual(UniformPurchase.objects.count(), 1); self.assertFalse(UniformType.objects.filter(name="Z").exists())

    def test_anonymous_redirected_and_mutations_are_post_only(self):
        p = self.purchase()
        self.client.logout()
        for u in self.URLS: self.assertEqual(self.client.get(reverse(u)).status_code, 302)
        self.login()
        for n in ("hrd_uniform_void", "hrd_uniform_mark"): self.assertEqual(self.client.get(reverse(n, args=[p.pk])).status_code, 405)

    def test_hub_card(self):
        self.assertContains(self.client.get("/hrd/"), "/hrd/uniforms/")


class RateTests(Base):
    def test_rate_by_gender_copied_to_row(self):
        self.buy(); self.buy(nik="002", quantity="3")
        a, b = UniformPurchase.objects.get(employee=self.e1), UniformPurchase.objects.get(employee=self.e2)
        self.assertEqual((a.gender, a.rate_amount, a.quantity, a.deduction_amount), ("L", Decimal(19000), 2, Decimal(38000)))
        self.assertEqual((b.gender, b.rate_amount, b.quantity, b.deduction_amount), ("P", Decimal(17000), 3, Decimal(51000)))
        self.assertEqual(a.created_by.username, "hrd")

    def test_new_rate_does_not_change_old_rows_and_applies_by_effective_date(self):
        old = self.purchase()
        UniformRate.objects.create(gender="L", amount=Decimal(25000), effective_from=TODAY, created_by=self.hrd)
        old.refresh_from_db(); self.assertEqual(old.rate_amount, Decimal(19000))
        self.buy(purchase_date=TODAY.isoformat(), size=self.M.pk, quantity="1")
        new = UniformPurchase.objects.get(size=self.M); self.assertEqual(new.rate_amount, Decimal(25000))
        self.assertEqual(services.uniform_rate_for("L", date(2020, 1, 1)).amount, Decimal(19000))   # tanggal sebelum tarif baru → tarif lama

    def test_no_rate_for_date_is_rejected(self):
        from apps.hr.models import Employee
        self.assertIsNone(services.uniform_rate_for("L", date(1999, 12, 31)))
        Employee.objects.create(nik="777", name="Senior", gender="L", department=self.d1, position=self.pos, join_date=date(1990, 1, 1))
        r = self.buy(nik="777", purchase_date="1999-12-31"); self.assertEqual(r.status_code, 200); self.assertContains(r, "Belum ada tarif")
        self.assertEqual(UniformPurchase.objects.count(), 0)

    def test_rate_rows_are_append_only(self):
        r = UniformRate.objects.first()
        r.amount = Decimal(1)
        with self.assertRaises(PermissionError): r.save()
        with self.assertRaises(PermissionError): r.delete()

    def test_add_rate_via_master_and_duplicate_date_rejected(self):
        post = {"action": "add_rate", "r-gender": "P", "r-amount": "18.000", "r-effective_from": TODAY.isoformat()}
        self.assertEqual(self.client.post(reverse("hrd_uniform_master"), post).status_code, 302)
        self.assertEqual(UniformRate.objects.get(gender="P", effective_from=TODAY).amount, Decimal(18000))
        self.assertEqual(self.client.post(reverse("hrd_uniform_master"), post).status_code, 200)     # tanggal sama → ditolak
        self.assertEqual(UniformRate.objects.filter(gender="P", effective_from=TODAY).count(), 1)
        self.assertTrue(AuditLog.objects.filter(action="uniform_master_rate_add").exists())


class PurchaseValidationTests(Base):
    def test_create_and_audit_without_sensitive_fields(self):
        self.assertEqual(self.buy().status_code, 302)
        a = self.last_audit("uniform_purchase"); self.assertEqual((a.after["employee"], a.after["size"], a.after["qty"]), ("001", "L", 2))
        self.assertNotIn("0001234567890", str(a.after))

    def test_inactive_unknown_and_bad_values_rejected(self):
        for bad in ({"nik": "009"}, {"nik": "999"}, {"quantity": "0"}, {"quantity": "51"}, {"quantity": "abc"},
                    {"purchase_date": (TODAY + timedelta(days=1)).isoformat()}, {"purchase_date": "2023-12-31"}, {"purchase_date": "bukan-tanggal"}, {"utype": "9999"}, {"size": ""}):
            self.assertEqual(self.buy(**bad).status_code, 200, bad)
        self.assertEqual(UniformPurchase.objects.count(), 0)

    def test_duplicate_rejected_but_allowed_after_void_and_for_other_size(self):
        self.buy(); self.assertEqual(self.buy().status_code, 200); self.assertEqual(UniformPurchase.objects.count(), 1)
        self.assertEqual(self.buy(size=self.M.pk).status_code, 302)                        # ukuran lain boleh
        services.void_uniform_purchase(UniformPurchase.objects.get(size=self.L).pk, "salah input", self.hrd)
        self.assertEqual(self.buy().status_code, 302)                                      # setelah batal boleh dicatat ulang
        self.assertEqual(UniformPurchase.objects.filter(voided_at__isnull=True).count(), 2)

    def test_inactive_type_and_size_not_selectable(self):
        self.type.is_active = False; self.type.save(); self.assertEqual(self.buy().status_code, 200)
        self.type.is_active = True; self.type.save(); self.L.is_active = False; self.L.save(); self.assertEqual(self.buy().status_code, 200)
        self.assertEqual(UniformPurchase.objects.count(), 0)

    def test_purchase_cannot_be_deleted(self):
        p = self.purchase()
        with self.assertRaises(PermissionError): p.delete()


class VoidAndMarkTests(Base):
    def test_void_requires_reason_and_removes_from_recap(self):
        p = self.purchase()
        r = self.client.post(reverse("hrd_uniform_void", args=[p.pk]), {"reason": "  "}); p.refresh_from_db(); self.assertFalse(p.is_void)
        self.client.post(reverse("hrd_uniform_void", args=[p.pk]), {"reason": "salah ukuran"}); p.refresh_from_db()
        self.assertTrue(p.is_void); self.assertEqual((p.void_reason, p.voided_by.username), ("salah ukuran", "hrd"))
        self.assertEqual(self.client.get(reverse("hrd_uniforms"), {"period": PER}).context["totals"]["n"], 0)
        self.assertEqual(self.client.get(reverse("hrd_uniforms"), {"period": PER, "batal": 1, "tab": "pembelian"}).context["page"].paginator.count, 1)   # tetap tampil bila diminta
        self.assertEqual(self.last_audit("uniform_void").after["reason"], "salah ukuran")

    def test_cannot_void_twice_or_after_deducted(self):
        p = self.purchase()
        services.mark_uniform_deducted(p.pk, PER, self.hrd)
        with self.assertRaises(ValueError): services.void_uniform_purchase(p.pk, "x", self.hrd)
        q = self.purchase(self.e2); services.void_uniform_purchase(q.pk, "x", self.hrd)
        with self.assertRaises(ValueError): services.void_uniform_purchase(q.pk, "lagi", self.hrd)
        with self.assertRaises(ValueError): services.mark_uniform_deducted(q.pk, PER, self.hrd)

    def test_mark_deducted_is_one_way_and_validates_period(self):
        p = self.purchase()
        for bad in ("", "2026-13", "26-10", "2000-01", "xxxx-xx"):
            self.client.post(reverse("hrd_uniform_mark", args=[p.pk]), {"period": bad}); p.refresh_from_db(); self.assertEqual(p.deduction_status, "belum", bad)
        self.client.post(reverse("hrd_uniform_mark", args=[p.pk]), {"period": PER}); p.refresh_from_db()
        self.assertEqual((p.deduction_status, p.deducted_period), ("sudah", PER)); self.assertTrue(self.last_audit("uniform_mark_deducted"))
        with self.assertRaises(ValueError): services.mark_uniform_deducted(p.pk, PER, self.hrd)    # tidak ada 'batal tandai'

    def test_detail_page_hides_actions_when_not_applicable(self):
        p = self.purchase(); r = self.client.get(reverse("hrd_uniform_detail", args=[p.pk])); self.assertContains(r, "Batalkan pembelian"); self.assertContains(r, "Tandai sudah dipotong")
        services.mark_uniform_deducted(p.pk, PER, self.hrd); r = self.client.get(reverse("hrd_uniform_detail", args=[p.pk])); self.assertNotContains(r, "Batalkan pembelian")

    def test_unknown_pk_is_404(self):
        for n in ("hrd_uniform_detail", "hrd_uniform_void", "hrd_uniform_mark"):
            m = self.client.post if n != "hrd_uniform_detail" else self.client.get
            self.assertEqual(m(reverse(n, args=[99999]), {"reason": "x", "period": PER}).status_code, 404, n)


class RecapTests(Base):
    def setUp(self):
        super().setUp()
        self.purchase(self.e1, quantity=2)                          # L Produksi  L-size  2 × 19.000 = 38.000
        self.purchase(self.e2, quantity=3, size=self.M)             # P Gudang    M-size  3 × 17.000 = 51.000
        self.purchase(self.e3, quantity=1)                          # P Produksi  L-size  1 × 17.000 = 17.000
        self.void = self.purchase(self.e1, quantity=5, size=self.M)  # batal: tak boleh masuk rekap
        services.void_uniform_purchase(self.void.pk, "salah", self.hrd)

    def get(self, **q): return self.client.get(reverse("hrd_uniforms"), {"period": PER, "tab": "pembelian", **q})

    def test_totals_exclude_void(self):
        t = self.get().context["totals"]; self.assertEqual((t["n"], t["pcs"], t["total"]), (3, 6, Decimal(106000)))
        self.assertEqual(t["belum"], Decimal(106000))
        services.mark_uniform_deducted(UniformPurchase.objects.get(employee=self.e2).pk, PER, self.hrd)
        self.assertEqual(self.get().context["totals"]["belum"], Decimal(55000))

    def test_size_by_gender_matrix(self):
        m = {(r["size"]): r for r in self.get(tab="stok").context["stock"]}
        self.assertEqual((m["L"]["L"], m["L"]["P"], m["L"]["keluar"]), (2, 1, 3)); self.assertEqual((m["M"]["L"], m["M"]["P"], m["M"]["keluar"]), (0, 3, 3))
        self.assertEqual([r["size"] for r in self.get(tab="stok").context["stock"]], ["M", "L"])        # urutan master (M=20, L=30), bukan abjad

    def test_gender_is_the_copy_not_current_employee_gender(self):
        self.e1.gender = "P"; self.e1.save()      # data karyawan berubah setelah pembelian
        m = {r["size"]: r for r in self.get(tab="stok").context["stock"]}; self.assertEqual(m["L"]["L"], 2)

    def test_per_department_and_per_employee(self):
        d = {r["employee__department__name"]: r for r in self.get().context["by_dept"]}
        self.assertEqual((d["Produksi"]["n"], d["Produksi"]["pcs"], d["Produksi"]["total"]), (2, 3, Decimal(55000))); self.assertEqual((d["Gudang"]["pcs"], d["Gudang"]["total"]), (3, Decimal(51000)))
        e = self.get().context["employees"]; self.assertEqual([x["employee__nik"] for x in e], ["002", "001", "003"])    # urut potongan terbesar
        self.assertEqual(self.get().context["employees_n"], 3)

    def test_filters(self):
        self.assertEqual(self.get(department=self.d2.pk).context["totals"]["n"], 1)
        self.assertEqual(self.get(size=self.M.pk).context["totals"]["n"], 1)
        self.assertEqual(self.get(status="sudah").context["totals"]["n"], 0)
        self.assertEqual(self.get(q="sari").context["totals"]["n"], 1)
        self.assertEqual(self.get(period="2001-01").context["totals"]["n"], 0)
        self.assertEqual(self.client.get(reverse("hrd_uniforms"), {"period": ""}).context["totals"]["n"], 3)         # kosong = semua periode

    def test_default_period_is_this_month_and_rupiah_format(self):
        r = self.client.get(reverse("hrd_uniforms")); self.assertEqual(r.context["f"]["period"], PER)
        self.assertContains(r, "Rp 106.000"); self.assertNotContains(r, ",00")

    def test_garbage_filters_do_not_500(self):
        for q in ("period=zzz", "department=abc", "utype=%00", "size=-1", "status=%27", "page=999", "export=zzz"):
            self.assertIn(self.client.get(reverse("hrd_uniforms") + "?" + q).status_code, (200, 400), q)

    def test_pagination_50(self):
        # banyak baris: pakai ukuran berbeda agar lolos constraint unik
        for i in range(60):
            z = UniformSize.objects.create(code=f"Z{i}", sort=900 + i); self.purchase(self.e1, size=z)
        r = self.get(); self.assertEqual(len(r.context["page"]), 50); self.assertEqual(r.context["page"].paginator.count, 63)


class ExportTests(Base):
    def setUp(self):
        super().setUp(); self.purchase(self.e1, quantity=2); self.purchase(self.e2, quantity=3, size=self.M)

    def sheet(self, kind, **q):
        r = self.client.get(reverse("hrd_uniforms"), {"period": PER, "export": kind, "format": "xlsx", **q}); self.assertEqual(r.status_code, 200)
        return list(load_workbook(io.BytesIO(r.content)).worksheets[0].iter_rows())

    def test_detail_export_money_cells_are_numeric(self):
        rows = self.sheet("detail"); head = [c.value for c in rows[0]]; self.assertEqual(head[:3], ["nik", "nama", "departemen"])
        i, j = head.index("tarif_per_satuan"), head.index("total_potongan")
        for r in rows[1:]: self.assertIsInstance(r[j].value, (int, float)); self.assertEqual(r[j].number_format, XLSX_RUPIAH_FORMAT); self.assertEqual(r[i].number_format, XLSX_RUPIAH_FORMAT)
        self.assertEqual(sum(r[j].value for r in rows[1:]), 38000 + 51000)

    def test_recap_export_matches_matrix(self):
        rows = [[c.value for c in r] for r in self.sheet("recap")]
        self.assertEqual(rows[0], ["jenis", "ukuran", "stok_awal", "masuk", "keluar_laki_laki", "keluar_perempuan", "keluar_total", "koreksi_kembali", "stok_akhir"])
        self.assertIn(["Seragam Kerja", "L", 0, 0, 2, 0, 2, 0, -2], rows); self.assertIn(["Seragam Kerja", "M", 0, 0, 0, 3, 3, 0, -3], rows)   # tanpa barang masuk → stok minus

    def test_export_has_no_bpjs_numbers_and_is_audited(self):
        for kind in ("detail", "recap"):
            body = self.client.get(reverse("hrd_uniforms"), {"period": PER, "export": kind}).content
            self.assertNotIn(b"0001234567890", body); self.assertNotIn(b"TK-9988776655", body)
        self.assertEqual(AuditLog.objects.filter(action="uniform_export").count(), 2)

    def test_detail_export_excludes_void_unless_requested(self):
        v = self.purchase(self.e3); services.void_uniform_purchase(v.pk, "salah", self.hrd)
        self.assertEqual(len(self.sheet("detail")) - 1, 2); rows = self.sheet("detail", batal=1); self.assertEqual(len(rows) - 1, 3)
        head = [c.value for c in rows[0]]; self.assertIn("batal", [r[head.index("status")].value for r in rows[1:]])

    def test_csv_export_plain_money(self):
        b = self.client.get(reverse("hrd_uniforms"), {"period": PER, "export": "detail"}).content.decode("utf-8-sig"); self.assertNotIn(".00", b); self.assertIn("38000", b)

    def test_formula_like_values_are_neutralised(self):
        self.purchase(self.e3, note="=HYPERLINK(\"x\")")
        b = self.client.get(reverse("hrd_uniforms"), {"period": PER, "export": "detail"}).content.decode("utf-8-sig"); self.assertIn("'=HYPERLINK", b)


class ImportTests(Base):
    def test_check_only_saves_nothing(self):
        r = self.upload(HDR + f"001,{D0},Seragam Kerja,L,2,\n", mode="check"); self.assertContains(r, "valid"); self.assertEqual(UniformPurchase.objects.count(), 0)

    def test_import_creates_with_rate_by_gender(self):
        self.upload(HDR + f"001,{D0},seragam kerja,l,2,baru\n002,{D0},Seragam Kerja,M,1,\n")
        a, b = UniformPurchase.objects.get(employee=self.e1), UniformPurchase.objects.get(employee=self.e2)
        self.assertEqual((a.deduction_amount, a.size.code, a.note), (Decimal(38000), "L", "baru")); self.assertEqual(b.deduction_amount, Decimal(17000))
        self.assertTrue(AuditLog.objects.filter(module="hrd", action="seragam_import").exists())

    def test_reimport_same_rows_is_rejected_not_duplicated(self):
        self.upload(HDR + f"001,{D0},Seragam Kerja,L,2,\n"); r = self.upload(HDR + f"001,{D0},Seragam Kerja,L,2,\n")
        self.assertEqual(UniformPurchase.objects.count(), 1); self.assertContains(r, "sudah tercatat")

    def test_all_or_nothing_and_bad_rows(self):
        for body in (f"001,{D0},Seragam Kerja,L,2,\n999,{D0},Seragam Kerja,L,2,\n", f"001,{D0},Jas Hujan,L,2,\n", f"001,{D0},Seragam Kerja,XXXL,2,\n",
                     f"009,{D0},Seragam Kerja,L,2,\n", f"001,{D0},Seragam Kerja,L,0,\n", f"001,2999-01-01,Seragam Kerja,L,1,\n", f"001,{D0},Seragam Kerja,L,1,=1+1\n",
                     f"001,{D0},Seragam Kerja,L,1,\n001,{D0},Seragam Kerja,L,1,\n"):
            self.upload(HDR + body); self.assertEqual(UniformPurchase.objects.count(), 0, body)

    def test_inactive_master_values_rejected(self):
        self.type.is_active = False; self.type.save(); self.upload(HDR + f"001,{D0},Seragam Kerja,L,1,\n"); self.assertEqual(UniformPurchase.objects.count(), 0)

    def test_xlsx_import(self):
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active; ws.append(["nik", "tanggal", "jenis", "ukuran", "jumlah", "catatan"]); ws.append(["001", D0.isoformat(), "Seragam Kerja", "L", 1, ""])
        out = io.BytesIO(); wb.save(out)
        self.client.post(reverse("hrd_uniform_import"), {"file": SimpleUploadedFile("u.xlsx", out.getvalue()), "mode": "import"}); self.assertEqual(UniformPurchase.objects.count(), 1)

    def test_template_download(self):
        r = self.client.get("/hrd/uniforms/import/template.csv"); self.assertEqual(r.status_code, 200); self.assertIn(b"ukuran", r.content)


class MasterTests(Base):
    def post(self, **kw): return self.client.post(reverse("hrd_uniform_master"), kw)

    def test_add_type_and_size_with_duplicates_rejected(self):
        self.post(action="add_type", name="Jaket"); self.post(action="add_type", name="jaket"); self.assertEqual(UniformType.objects.filter(name__iexact="jaket").count(), 1)
        self.post(action="add_size", code="4xl"); self.post(action="add_size", code="4XL"); self.assertEqual(UniformSize.objects.filter(code="4XL").count(), 1)
        self.assertEqual(UniformSize.objects.get(code="4XL").sort, 70)     # otomatis di urutan paling akhir
        for bad in ({"action": "add_type", "name": " "}, {"action": "add_size", "code": "x" * 11}, {"action": "nope"}): self.post(**bad)
        self.assertEqual(UniformType.objects.count(), 2)

    def test_toggle_keeps_history(self):
        p = self.purchase(); self.post(action="toggle_type", id=self.type.pk); self.type.refresh_from_db(); self.assertFalse(self.type.is_active)
        p.refresh_from_db(); self.assertEqual(p.utype_id, self.type.pk)        # pembelian lama tetap
        self.post(action="toggle_type", id=self.type.pk); self.type.refresh_from_db(); self.assertTrue(self.type.is_active)
        self.assertEqual(AuditLog.objects.filter(action="uniform_master_type_toggle").count(), 2)

    def test_page_shows_current_rates(self):
        r = self.client.get(reverse("hrd_uniform_master")); self.assertContains(r, "Rp 19.000"); self.assertContains(r, "Rp 17.000"); self.assertNotContains(r, ",00")
