"""Sisa P4 (putaran 23): sel XLSX rupiah = angka ber-format, ringkasan per departemen, URL submenu K/TK, sidebar BPJS."""
import io
from decimal import Decimal
from django.urls import reverse
from openpyxl import load_workbook
from apps.core import tabular
from apps.core.money import XLSX_RUPIAH_FORMAT
from .models import BpjsDeduction
from .test_base import HrdBase

P = "2026-10"


class Base(HrdBase):
    def setUp(self):
        self.login()
        self.mk(self.e1, "kes", 150000, 600000)       # Produksi
        self.mk(self.e1, "tk", 20000, 30000)          # Produksi
        self.mk(self.e2, "kes", 1250000, 0)           # Gudang
        self.mk(self.e3, "tk", 5000, 7500)           # Produksi

    def mk(self, e, scheme, emp, er, period=P):
        return BpjsDeduction.objects.create(employee=e, scheme=scheme, period=period, employee_amount=Decimal(emp), employer_amount=Decimal(er), created_by=self.hrd)

    def get(self, path="", **q):
        return self.client.get(reverse("hrd_bpjs_deductions") + path, {"period": P, **q})


class XlsxMoneyTests(Base):
    def sheet(self, **q):
        r = self.get(export=1, format="xlsx", **q); self.assertEqual(r.status_code, 200)
        return list(load_workbook(io.BytesIO(r.content)).worksheets[0].iter_rows())

    def test_money_cells_are_numbers_with_rupiah_format(self):
        rows = self.sheet(); head = [c.value for c in rows[0]]
        i_emp, i_er = head.index("porsi_karyawan"), head.index("porsi_perusahaan")
        for r in rows[1:]:
            for i in (i_emp, i_er):
                self.assertIsInstance(r[i].value, (int, float), r[i].value); self.assertEqual(r[i].number_format, XLSX_RUPIAH_FORMAT)
        self.assertEqual(sum(r[i_emp].value for r in rows[1:]), 150000 + 20000 + 1250000 + 5000)   # bisa dijumlah (bukan teks)

    def test_other_columns_stay_text_and_unformatted(self):
        rows = self.sheet(); head = [c.value for c in rows[0]]
        r = rows[1]; self.assertIsInstance(r[head.index("nik")].value, str); self.assertNotEqual(r[head.index("nik")].number_format, XLSX_RUPIAH_FORMAT)
        self.assertIsInstance(r[head.index("periode")].value, str)

    def test_csv_money_is_plain_number_without_cents(self):
        body = self.get(export=1).content.decode("utf-8-sig")
        self.assertNotIn(".00", body); self.assertNotIn(",00", body); self.assertIn("1250000", body)
        self.assertNotIn("Rp", body)  # CSV: angka polos; format rupiah hanya di tampilan/XLSX

    def test_csv_header_names_money_columns(self):
        lines = self.get(export=1).content.decode("utf-8-sig").splitlines()
        self.assertEqual(lines[0].split(",")[5:7], ["porsi_karyawan", "porsi_perusahaan"]); self.assertEqual(len(lines) - 1, 4)

    def test_export_respects_scheme_filter_and_is_audited(self):
        rows = self.sheet(scheme="tk"); self.assertEqual(len(rows) - 1, 2)
        a = self.last_audit("bpjs_deduction_export"); self.assertEqual(a.after["scheme"], "tk")


class TabularUnitTests(HrdBase):
    def test_money_cols_by_name_and_index(self):
        for mc in (("jumlah",), (1,)):
            b, _ = tabular.to_bytes(["nama", "jumlah"], [["a", Decimal("1500000.00")], ["b", Decimal("2.50")]], "xlsx", money_cols=mc)
            ws = load_workbook(io.BytesIO(b)).worksheets[0]
            self.assertEqual([ws.cell(r, 2).value for r in (2, 3)], [1500000, 2.5]); self.assertEqual(ws.cell(2, 2).number_format, XLSX_RUPIAH_FORMAT)

    def test_bool_and_text_never_treated_as_money(self):
        b, _ = tabular.to_bytes(["x", "y", "z"], [[True, "=1+1", "12"]], "xlsx", money_cols=(0, 1, 2))
        ws = load_workbook(io.BytesIO(b)).worksheets[0]
        self.assertEqual(ws.cell(2, 1).value, "ya"); self.assertEqual(ws.cell(2, 2).value, "'=1+1"); self.assertEqual(ws.cell(2, 3).value, "12")

    def test_num_cols_are_plain_numbers_without_money_format(self):
        b, _ = tabular.to_bytes(["nama", "pcs"], [["a", 3], ["b", Decimal("4")]], "xlsx", num_cols=("pcs",))
        ws = load_workbook(io.BytesIO(b)).worksheets[0]
        self.assertEqual([ws.cell(r, 2).value for r in (2, 3)], [3, 4]); self.assertNotEqual(ws.cell(2, 2).number_format, XLSX_RUPIAH_FORMAT)
        c, _ = tabular.to_bytes(["nama", "pcs"], [["a", 3]], "csv", num_cols=(1,)); self.assertIn("a,3", c.decode("utf-8-sig"))

    def test_default_behavior_unchanged_without_money_cols(self):
        b, _ = tabular.to_bytes(["a"], [[Decimal("1.50")]], "csv"); self.assertIn("1.50", b.decode("utf-8-sig"))


class DepartmentSummaryTests(Base):
    def rows(self, **q):
        r = self.get(**q); self.assertEqual(r.status_code, 200); return {d["employee__department__name"]: d for d in r.context["by_dept"]}, r

    def test_per_department_totals(self):
        d, r = self.rows()
        self.assertEqual((d["Produksi"]["n"], d["Produksi"]["emp"], d["Produksi"]["er"]), (3, Decimal(175000), Decimal(637500)))
        self.assertEqual((d["Gudang"]["n"], d["Gudang"]["emp"]), (1, Decimal(1250000)))
        self.assertEqual(sum(x["emp"] for x in d.values()), r.context["totals"]["emp"])      # jumlah departemen = total
        self.assertContains(r, 'id="by-dept"'); self.assertContains(r, "Rp 1.250.000"); self.assertNotContains(r, ",00")

    def test_follows_filters(self):
        d, _ = self.rows(scheme="tk"); self.assertEqual((d["Produksi"]["n"], d["Produksi"]["emp"]), (2, Decimal(25000))); self.assertNotIn("Gudang", d)
        d, _ = self.rows(department=self.d2.pk); self.assertEqual(list(d), ["Gudang"])
        d, _ = self.rows(period="2026-09"); self.assertEqual(d, {})

    def test_empty_period_has_no_table(self):
        self.assertNotContains(self.get(period="2025-01"), 'id="by-dept"')

    def test_hidden_in_missing_anomaly_view(self):
        self.assertNotContains(self.get(anomaly="belum"), 'id="by-dept"')


class SchemeUrlTests(Base):
    def test_scheme_urls_preselect_and_filter(self):
        for name, sc, n in (("hrd_bpjs_deductions_kes", "kes", 2), ("hrd_bpjs_deductions_tk", "tk", 2)):
            r = self.client.get(reverse(name), {"period": P}); self.assertEqual(r.status_code, 200)
            self.assertEqual(r.context["f"]["scheme"], sc); self.assertEqual(len(r.context["page"]), n)
            self.assertTrue(all(d.scheme == sc for d in r.context["page"]))

    def test_query_param_overrides_url_default_including_all(self):
        r = self.client.get(reverse("hrd_bpjs_deductions_kes"), {"period": P, "scheme": ""}); self.assertEqual(len(r.context["page"]), 4)
        r = self.client.get(reverse("hrd_bpjs_deductions_kes"), {"period": P, "scheme": "tk"}); self.assertEqual({d.scheme for d in r.context["page"]}, {"tk"})

    def test_export_from_scheme_page_keeps_scheme(self):
        r = self.client.get(reverse("hrd_bpjs_deductions_kes"), {"period": P}); self.assertContains(r, "scheme=kes&export=1")
        r = self.client.get(reverse("hrd_bpjs_deductions_kes"), {"period": P, "export": 1}); body = r.content.decode("utf-8-sig")
        self.assertNotIn("Ketenagakerjaan", body); self.assertEqual(len(body.splitlines()) - 1, 2)

    def test_title_names_the_program(self):
        self.assertContains(self.client.get(reverse("hrd_bpjs_deductions_kes")), "Kesehatan (K)")
        self.assertContains(self.client.get(reverse("hrd_bpjs_deductions_tk")), "Ketenagakerjaan (TK)")

    def test_rbac_and_anonymous(self):
        for who, code in (("hrd", 200), ("su", 200), ("adm", 403), ("poli", 403)):
            self.login(who)
            for n in ("hrd_bpjs_deductions_kes", "hrd_bpjs_deductions_tk"): self.assertEqual(self.client.get(reverse(n)).status_code, code, (who, n))
        self.client.logout()
        for n in ("hrd_bpjs_deductions_kes", "hrd_bpjs_deductions_tk"): self.assertEqual(self.client.get(reverse(n)).status_code, 302)

    def test_new_and_import_not_swallowed_by_scheme_routes(self):
        self.assertEqual(self.client.get(reverse("hrd_bpjs_deduction_new")).status_code, 200)
        self.assertEqual(self.client.get("/hrd/bpjs/deductions/import/").status_code, 200)


class SidebarTests(HrdBase):
    def nav(self, who, path):
        self.login(who); r = self.client.get(path); self.assertEqual(r.status_code, 200)
        return [i for g in r.context["nav_groups"] for i in g["items"]], r

    def test_hrd_and_superadmin_see_bpjs_with_submenu(self):
        for who in ("hrd", "su"):
            items, _ = self.nav(who, "/hrd/")
            labels = [i["label"] for i in items]; self.assertIn("BPJS", labels)
            subs = {i["label"]: i["url"] for i in items if i["sub"] and "/hrd/bpjs/" in i["url"]}
            self.assertEqual(subs["Kesehatan (K)"], "/hrd/bpjs/deductions/kes/"); self.assertEqual(subs["Ketenagakerjaan (TK)"], "/hrd/bpjs/deductions/tk/")

    def test_other_roles_do_not_see_bpjs_menu(self):
        for who, path in (("adm", "/"), ("poli", "/")):
            items, _ = self.nav(who, path); self.assertFalse([i for i in items if "/hrd/bpjs" in i["url"]], who)

    def test_active_item_follows_page(self):
        for path, label in (("/hrd/bpjs/deductions/kes/", "Kesehatan (K)"), ("/hrd/bpjs/deductions/tk/", "Ketenagakerjaan (TK)"), ("/hrd/bpjs/deductions/", "Semua Potongan"),
                            ("/hrd/bpjs/", "BPJS"), ("/hrd/", "Operasional HRD"), ("/hrd/aids/", "Bantuan (rekap)"), ("/hrd/catering/", "Katering (rekap)"), ("/hrd/uniforms/", "Seragam")):
            items, r = self.nav("hrd", path); act = [i["label"] for i in items if i["active"]]
            self.assertEqual(act, [label], path); self.assertEqual(r.context["crumb"], label)

    def test_status_detail_page_keeps_bpjs_active(self):
        items, _ = self.nav("hrd", f"/hrd/bpjs/{self.e1.pk}/"); self.assertEqual([i["label"] for i in items if i["active"]], ["BPJS"])


class OperasionalMenuTests(HrdBase):
    """Putaran 28: Operasional HRD = menu + submenu (Seragam, Katering, dst) dengan dropdown buka/tutup berikon."""
    SUBS = ["Bantuan (rekap)", "Cuti Hamil", "Pekerja Harian Proyek", "Katering (rekap)", "Surat Peringatan", "Seragam", "Karyawan Keluar"]

    def tree(self, who, path):
        self.login(who); r = self.client.get(path); self.assertEqual(r.status_code, 200)
        return {n["label"]: n for g in r.context["nav_groups"] for n in g["tree"]}, r

    def test_operasional_has_children_and_seragam_moved_inside(self):
        for who in ("hrd", "su"):
            t, _ = self.tree(who, "/hrd/")
            self.assertEqual([c["label"] for c in t["Operasional HRD"]["children"]], self.SUBS)
            self.assertNotIn("Seragam", t)  # tidak lagi menu tingkat atas
            self.assertEqual(t["Operasional HRD"]["url"], "/hrd/")

    def test_parents_with_submenu_are_all_dropdowns(self):
        t, _ = self.tree("hrd", "/")
        for parent in ("Operasional HRD", "BPJS", "Semua Pengajuan"): self.assertTrue(t[parent]["children"], parent)
        self.assertFalse(t["Dashboard"]["children"])

    def test_closed_by_default_open_when_page_inside(self):
        t, r = self.tree("hrd", "/")
        self.assertFalse(t["Operasional HRD"]["open"]); self.assertContains(r, 'aria-controls="sub-operasional-hrd"'); self.assertContains(r, 'aria-expanded="false"')
        t, r = self.tree("hrd", "/hrd/catering/"); self.assertTrue(t["Operasional HRD"]["open"])
        self.assertEqual([c["label"] for c in t["Operasional HRD"]["children"] if c["active"]], ["Katering (rekap)"]); self.assertFalse(t["BPJS"]["open"])
        t, _ = self.tree("hrd", "/hrd/uniforms/stock/"); self.assertTrue(t["Operasional HRD"]["open"])  # halaman turunan tetap membuka induknya

    def test_toggle_has_icon_and_label_and_other_roles_have_no_ops_menu(self):
        _, r = self.tree("hrd", "/hrd/"); self.assertContains(r, 'class="ntog"'); self.assertContains(r, "Buka atau tutup submenu Operasional HRD"); self.assertContains(r, 'href="#i-chev"')
        for who in ("adm", "poli"):
            t, _ = self.tree(who, "/"); self.assertNotIn("Operasional HRD", t, who)
