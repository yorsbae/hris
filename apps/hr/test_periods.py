"""Tutup periode (tutup buku) dengan rentang tanggal bebas (putaran 59): checklist, penguncian absensi & tagihan Poli, riwayat, snapshot XLSX, buka kembali hanya Superadmin."""
import io
import tempfile
from datetime import date, time

from django.test import TestCase, override_settings
from openpyxl import load_workbook

from apps.core.models import AuditLog, Notification, Role, User
from apps.hr import period_export, periods
from apps.hr.models import AttendanceDaily, Department, Employee, PeriodClose, Shift

PW = "Pw-12345-abc!"
TMP = tempfile.mkdtemp()


@override_settings(MEDIA_ROOT=TMP)
class PeriodTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.d = Department.objects.create(code="PR", name="Produksi")
        cls.hrd = User.objects.create_user("hrd9", password=PW, role=Role.HRD); cls.su = User.objects.create_user("su9", password=PW, role=Role.SUPERADMIN)
        cls.pay = User.objects.create_user("pay9", password=PW, role=Role.PAYROLL)
        cls.adm = User.objects.create_user("adm9", password=PW, role=Role.DEPT_ADMIN, department=cls.d)
        cls.sh = Shift.objects.create(name="Pagi9", code="PAGI9", start=time(7), end=time(15))
        cls.e = Employee.objects.create(nik="P1", name="Pee", gender="L", department=cls.d, join_date=date(2020, 1, 1), shift=cls.sh)

    def login(self, u): self.client.login(username=u, password=PW)

    def day(self, d, status="hadir", **k): return AttendanceDaily.objects.create(employee=self.e, department=self.d, date=d, status=status, **k)

    def close(self, a=date(2026, 9, 26), b=date(2026, 10, 15), label="2026-10", **k):
        return self.client.post("/absensi/periode/", {"action": "close", "start": str(a), "end": str(b), "label": label, **k})

    # ---- logika rentang
    def test_segments_skip_closed_range_not_calendar_month(self):
        PeriodClose.objects.create(period="2026-10", start_date=date(2026, 9, 26), end_date=date(2026, 10, 25))
        self.assertEqual(periods.open_segments(date(2026, 9, 20), date(2026, 11, 5)), [(date(2026, 9, 20), date(2026, 9, 25)), (date(2026, 10, 26), date(2026, 11, 5))])
        self.assertEqual(periods.open_segments(date(2026, 10, 1), date(2026, 10, 3)), [])
        self.assertEqual(periods.open_segments(date(2026, 11, 1), date(2026, 11, 3)), [(date(2026, 11, 1), date(2026, 11, 3))])
        self.assertTrue(periods.is_closed(date(2026, 9, 26))); self.assertTrue(periods.is_closed(date(2026, 10, 25))); self.assertFalse(periods.is_closed(date(2026, 10, 26)))

    def test_reopened_range_does_not_lock(self):
        PeriodClose.objects.create(period="2026-10", start_date=date(2026, 9, 26), end_date=date(2026, 10, 25), closed=False)
        self.assertEqual(periods.open_segments(date(2026, 10, 1), date(2026, 10, 3)), [(date(2026, 10, 1), date(2026, 10, 3))])

    def test_closed_range_is_not_recomputed(self):
        from apps.hr import attendance_calc as calc
        today = date(2026, 10, 20)
        calc.process(date(2026, 9, 28), date(2026, 9, 30), [self.e.pk], today); self.assertEqual(AttendanceDaily.objects.filter(date__month=9).count(), 3)
        PeriodClose.objects.create(period="2026-10", start_date=date(2026, 9, 26), end_date=date(2026, 10, 5)); AttendanceDaily.objects.all().delete()
        self.assertEqual(calc.process(date(2026, 9, 28), date(2026, 9, 30), [self.e.pk], today)["rows"], 0); self.assertEqual(AttendanceDaily.objects.count(), 0)
        calc.process(date(2026, 10, 6), date(2026, 10, 8), [self.e.pk], today); self.assertEqual(AttendanceDaily.objects.count(), 3)   # di luar rentang tetap jalan

    # ---- tutup lewat halaman
    def test_close_builds_history_row_and_xlsx_with_all_sheets(self):
        self.day(date(2026, 9, 28), overtime_minutes=90, overtime_approved_minutes=90); self.day(date(2026, 9, 29), standby_minutes=30)
        self.login("hrd9"); self.assertEqual(self.close().status_code, 302)
        row = PeriodClose.objects.get(); self.assertEqual((row.start_date, row.end_date, row.period, row.closed, row.closed_by), (date(2026, 9, 26), date(2026, 10, 15), "2026-10", True, self.hrd))
        self.assertTrue(row.export_file); self.assertEqual(list(row.export_counts), list(period_export.SHEETS))
        wb = load_workbook(row.export_file.open("rb")); self.assertEqual(wb.sheetnames, list(period_export.SHEETS))
        att = list(wb["Rekap Absensi"].iter_rows(values_only=True)); self.assertEqual(att[0][0], "NIK"); self.assertEqual(att[1][0], "P1"); self.assertEqual(att[1][5], 2)
        ot = list(wb["Lembur & Standby"].iter_rows(values_only=True)); self.assertEqual({r[4] for r in ot[1:]}, {"Lembur", "Stand By"})
        a = AuditLog.objects.get(action="period_close"); self.assertEqual(a.after["start"], "2026-09-26")

    def test_overlap_and_duplicate_label_are_refused_but_adjacent_is_fine(self):
        self.login("hrd9"); self.close(date(2026, 9, 26), date(2026, 10, 15), "2026-10", force="1")
        self.close(date(2026, 10, 10), date(2026, 10, 18), "2026-11", force="1"); self.assertEqual(PeriodClose.objects.count(), 1)         # tumpang tindih
        self.close(date(2026, 10, 16), date(2026, 10, 18), "2026-10", force="1"); self.assertEqual(PeriodClose.objects.count(), 1)         # label gaji sama
        self.close(date(2026, 10, 16), date(2026, 10, 18), "2026-11", force="1"); self.assertEqual(PeriodClose.objects.count(), 2)         # bersebelahan boleh

    def test_invalid_ranges_are_refused(self):
        self.login("hrd9")
        self.close(date(2026, 10, 15), date(2026, 10, 1), force="1"); self.close(date(2026, 1, 1), date(2026, 10, 1), force="1")             # terbalik, > 93 hari
        self.close(date(2026, 9, 1), date(2026, 9, 5), "bukan", force="1"); self.close(date(2099, 1, 1), date(2099, 1, 5), "2099-01", force="1")   # label salah, belum lewat
        self.assertEqual(PeriodClose.objects.count(), 0)

    def test_blockers_stop_closing_unless_forced(self):
        self.day(date(2026, 9, 28), "tanpa_scan"); self.login("hrd9")
        self.close(); self.assertEqual(PeriodClose.objects.count(), 0)
        self.close(force="1"); row = PeriodClose.objects.get(); self.assertTrue(row.closed); self.assertTrue(AuditLog.objects.get(action="period_close").after["forced"])

    def test_reopen_is_superadmin_only_keeps_history_and_allows_reclose(self):
        self.day(date(2026, 9, 28)); self.login("hrd9"); self.close(); row = PeriodClose.objects.get()
        self.client.post("/absensi/periode/", {"action": "reopen", "id": row.pk}); self.assertTrue(PeriodClose.objects.get().closed)          # HRD tidak boleh
        self.login("su9"); self.client.post("/absensi/periode/", {"action": "reopen", "id": row.pk}); row.refresh_from_db()
        self.assertFalse(row.closed); self.assertEqual(row.reopened_by, self.su); self.assertTrue(row.export_file)                          # berkas lama tetap sebagai riwayat
        self.assertTrue(AuditLog.objects.filter(action="period_reopen").exists())
        self.login("hrd9"); self.close(); self.assertEqual(PeriodClose.objects.count(), 2); self.assertEqual(PeriodClose.objects.filter(closed=True).count(), 1)   # riwayat bertambah

    def test_payroll_is_notified_on_close_and_reopen(self):
        self.day(date(2026, 9, 28)); self.login("hrd9"); self.close(); row = PeriodClose.objects.get()
        self.assertTrue(Notification.objects.filter(user=self.pay, title__contains="ditutup").exists())
        self.login("su9"); self.client.post("/absensi/periode/", {"action": "reopen", "id": row.pk}); self.assertTrue(Notification.objects.filter(user=self.pay, title__contains="DIBUKA KEMBALI").exists())

    def test_rebuild_refreshes_file(self):
        self.day(date(2026, 9, 28)); self.login("hrd9"); self.close(); row = PeriodClose.objects.get(); old = row.export_file.name
        self.client.post("/absensi/periode/", {"action": "rebuild", "id": row.pk}); row.refresh_from_db(); self.assertNotEqual(row.export_file.name, old); self.assertTrue(row.export_file)

    # ---- akses
    def test_page_and_download_access(self):
        self.day(date(2026, 9, 28)); self.login("hrd9"); self.close(); row = PeriodClose.objects.get()
        self.login("hrd9"); self.assertEqual(self.client.get("/absensi/periode/").status_code, 200); self.assertEqual(self.client.get(f"/absensi/periode/{row.pk}/unduh/").status_code, 200)
        self.assertTrue(AuditLog.objects.filter(action="period_export_download").exists())
        self.login("pay9"); self.assertEqual(self.client.get(f"/absensi/periode/{row.pk}/unduh/").status_code, 200); self.assertEqual(self.client.get("/absensi/periode/").status_code, 403)
        self.login("adm9"); self.assertEqual(self.client.get("/absensi/periode/").status_code, 403); self.assertEqual(self.client.get(f"/absensi/periode/{row.pk}/unduh/").status_code, 403)
        self.client.logout(); self.assertIn(self.client.get(f"/absensi/periode/{row.pk}/unduh/").status_code, (302, 401))

    def test_checklist_preview_for_free_range(self):
        self.login("hrd9"); r = self.client.get("/absensi/periode/", {"start": "2026-09-26", "end": "2026-10-15", "label": "2026-10"}); self.assertContains(r, "Checklist 26-09-2026 s/d 15-10-2026")

    def test_export_cells_are_formula_safe(self):
        self.e.name = "=HYPERLINK(1)"; self.e.save(); self.day(date(2026, 9, 28))
        data, _ = period_export.build(date(2026, 9, 26), date(2026, 10, 15), "2026-10")
        ws = load_workbook(io.BytesIO(data))["Rekap Absensi"]; self.assertEqual(ws.cell(row=2, column=2).value, "'=HYPERLINK(1)")


class PoliBillLockTests(PeriodTests):
    def test_bill_in_closed_range_is_refused_outside_is_allowed(self):
        from apps.poli import services
        from apps.poli.models import Partner
        PeriodClose.objects.create(period="2026-10", start_date=date(2026, 9, 26), end_date=date(2026, 10, 25))
        p = Partner.objects.create(name="RS Uji")
        with self.assertRaisesMessage(ValueError, "sudah ditutup"):
            services.create_partner_bill(self.hrd, p, "INV-1", date(2026, 10, 30), self.e, date(2026, 10, 10), "rawat_jalan", "", None, "non_bpjs", 100000)
        services.create_partner_bill(self.hrd, p, "INV-2", date(2026, 10, 30), self.e, date(2026, 10, 26), "rawat_jalan", "", None, "non_bpjs", 100000)
