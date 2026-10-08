"""Putaran 18: impor/ekspor CSV & XLSX (tabular, master, rotasi, karyawan, Poli, ekspor daftar)."""
import io
from datetime import date, time
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from apps.core import tabular
from apps.core.models import Role, User
from apps.poli.models import Diagnosis, Medicine
from .models import Department, Employee, Position, Shift, ShiftGroup, ShiftRotation

PW = "Sandi-Uji-12345"


def xlsx_bytes(rows):
    wb = Workbook(); ws = wb.active
    for r in rows: ws.append(r)
    b = io.BytesIO(); wb.save(b); return b.getvalue()


class TabularTests(TestCase):
    def test_csv_and_xlsx_roundtrip_and_formula_guard(self):
        for kind in ("csv", "xlsx"):
            b, _ = tabular.to_bytes(["a", "b"], [["=1+1", "x"], ["Budi", date(2026, 1, 2)]], kind)
            rows = tabular.read_table(b, f"f.{kind}", ["a"])
            self.assertEqual(rows[0]["a"], "'=1+1"); self.assertEqual(rows[1], {"a": "Budi", "b": "2026-01-02"})

    def test_semicolon_csv_missing_column_and_limits(self):
        self.assertEqual(tabular.read_table(b"a;b\n1;2\n", "x.csv", ["a"]), [{"a": "1", "b": "2"}])
        with self.assertRaises(ValueError): tabular.read_table(b"a;b\n1;2\n", "x.csv", ["zzz"])
        with self.assertRaises(ValueError): tabular.read_table(b"PK\x03\x04bukan-xlsx", "x.xlsx")
        with self.assertRaises(ValueError): tabular.read_table(b"x", "x.xls")
        with self.assertRaises(ValueError): tabular.read_table(b"a\n" + b"1\n" * 5001, "x.csv", ["a"])


class BulkBase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.d = Department.objects.create(code="PRD", name="Produksi")
        cls.hrd = User.objects.create_user("hrd", password=PW, role=Role.HRD)
        cls.poli = User.objects.create_user("poli", password=PW, role=Role.POLI)
        cls.adm = User.objects.create_user("adm", password=PW, role=Role.DEPT_ADMIN, department=cls.d)
        cls.sup = User.objects.create_user("sup", password=PW, role=Role.SUPERADMIN)

    def login(self, u): self.client.logout(); self.assertTrue(self.client.login(username=u, password=PW))

    def post(self, url, rows, mode="import", name="d.xlsx"):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return self.client.post(url, {"file": SimpleUploadedFile(name, xlsx_bytes(rows)), "mode": mode})


class MasterImportTests(BulkBase):
    def test_shift_import_upsert_and_check_mode(self):
        self.login("hrd"); rows = [["code", "name", "start", "end", "crosses_midnight", "is_gs", "active"],
                                   ["GS-16", "General Shift 16", "08:00", "16:00", "tidak", "ya", "ya"], ["MALAM", "Shift Malam", "23:00", "07:00", "ya", "tidak", "ya"]]
        self.post("/master/shift/import/", rows, mode="check"); self.assertEqual(Shift.objects.count(), 0)  # periksa saja: tidak menyimpan
        r = self.post("/master/shift/import/", rows); self.assertContains(r, "2 ditambahkan"); self.assertEqual(Shift.objects.count(), 2)
        self.assertTrue(Shift.objects.get(code="GS-16").is_gs); self.assertTrue(Shift.objects.get(code="MALAM").crosses_midnight)
        rows[1][3] = "15:00"; r = self.post("/master/shift/import/", rows); self.assertContains(r, "1 diperbarui".replace("1 diperbarui", "2 diperbarui"))
        self.assertEqual(Shift.objects.get(code="GS-16").end, time(15))

    def test_all_or_nothing_on_bad_row(self):
        self.login("hrd")
        r = self.post("/master/shift/import/", [["code", "name", "start", "end"], ["OK1", "Ok", "08:00", "16:00"], ["BAD", "Bad", "08:00", "08:00"]])
        self.assertContains(r, "Tidak ada yang disimpan"); self.assertEqual(Shift.objects.count(), 0)

    def test_department_parent_in_same_file_and_position(self):
        self.login("hrd")
        self.post("/master/department/import/", [["code", "name", "parent_code"], ["ANAK", "Anak", "INDUK"], ["INDUK", "Induk", ""]])  # induk setelah anak → ditolak
        self.assertFalse(Department.objects.filter(code="ANAK").exists())
        self.post("/master/department/import/", [["code", "name", "parent_code"], ["INDUK", "Induk", ""], ["ANAK", "Anak", "INDUK"]])
        self.assertEqual(Department.objects.get(code="ANAK").parent.code, "INDUK")
        self.post("/master/position/import/", [["name", "level"], ["Operator", "1"], ["operator", "2"]]); self.assertEqual(Position.objects.count(), 0)  # dobel (huruf beda)
        self.post("/master/position/import/", [["name", "level"], ["Operator", "1"]]); self.assertEqual(Position.objects.count(), 1)

    def test_rotation_import_creates_groups_by_code_and_validates(self):
        self.login("hrd"); pagi = Shift.objects.create(code="PAGI", name="Pagi", start=time(7), end=time(15)); malam = Shift.objects.create(code="MALAM", name="Malam", start=time(23), end=time(7), crosses_midnight=True)
        r = self.post("/master/rotasi/import/", [["group", "weekday", "shift_code"], ["A7", "senin", "PAGI"], ["A7", "selasa", "MALAM"], ["A7", "rabu", ""], ["A7_pack", "senin", "PAGI"]])
        self.assertContains(r, "4 ditambahkan")
        self.assertEqual(ShiftGroup.objects.get(code="A7").pattern, "3_SHIFT"); self.assertEqual(ShiftGroup.objects.get(code="A7_pack").pattern, "2_SHIFT")
        self.assertIsNone(ShiftRotation.objects.get(group__code="A7", weekday=2).shift)  # kosong = libur
        for bad in ([["group", "weekday", "shift_code"], ["A7_pack", "selasa", "MALAM"]], [["group", "weekday", "shift_code"], ["A", "senin", "PAGI"]], [["group", "weekday", "shift_code"], ["B7", "hari", "PAGI"]]):
            self.assertContains(self.post("/master/rotasi/import/", bad), "Tidak ada yang disimpan")
        self.assertFalse(ShiftGroup.objects.filter(code__in=("B7",)).exists())

    def test_import_requires_hrd(self):
        for u, code in (("poli", 403), ("adm", 403)):
            self.login(u); self.assertEqual(self.client.get("/master/shift/import/").status_code, code)
        self.client.logout(); self.assertEqual(self.client.get("/master/shift/import/").status_code, 302)

    def test_master_export_csv_and_xlsx(self):
        Shift.objects.create(code="GS-16", name="GS 16", start=time(8), end=time(16), is_gs=True); self.login("hrd")
        r = self.client.get("/master/shift/export/?format=xlsx"); ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual([c.value for c in ws[2]][:4], ["GS-16", "GS 16", "08:00", "16:00"])
        self.assertIn("GS-16", self.client.get("/master/shift/export/").content.decode("utf-8-sig"))
        self.login("poli"); self.assertEqual(self.client.get("/master/shift/export/").status_code, 403)


class EmployeeImportExportTests(BulkBase):
    def test_xlsx_import_with_group_and_gs_short_then_export_scope(self):
        Position.objects.create(name="Operator"); g = Shift.objects.create(code="GS-16", name="GS 16", start=time(8), end=time(16), is_gs=True)
        ShiftGroup.objects.create(code="A7", pattern="3_SHIFT"); self.login("hrd")
        hdr = ["nik", "name", "gender", "join_date", "department_code", "position", "shift", "shift_group", "gs_short"]
        r = self.post("/employees/import/", [hdr, ["E1", "Satu", "L", "2024-01-15", "PRD", "Operator", "GS-16", "", "12"], ["E2", "Dua", "P", "2024-02-01", "PRD", "Operator", "", "A7", ""]])
        self.assertContains(r, "2 karyawan diimpor")
        e1, e2 = Employee.objects.get(nik="E1"), Employee.objects.get(nik="E2")
        self.assertEqual((e1.shift_id, e1.gs_short, e2.shift_group.code), (g.pk, "12", "A7"))
        bad = self.post("/employees/import/", [hdr, ["E3", "Tiga", "L", "2024-01-15", "PRD", "", "", "Z9", ""]]); self.assertContains(bad, "Kelompok shift")
        x = load_workbook(io.BytesIO(self.client.get("/api/employees/export/?format=xlsx").content)).active
        self.assertEqual({r[0].value for r in list(x.rows)[1:]}, {"E1", "E2"}); self.assertEqual([c.value for c in x[1]][:2], ["nik", "name"])
        self.login("adm"); self.assertEqual(self.client.get("/api/employees/export/").status_code, 200)
        self.login("poli"); self.assertEqual(self.client.get("/api/employees/export/").status_code, 403)

    def test_dept_admin_export_limited_to_own_department(self):
        other = Department.objects.create(code="QC", name="QC")
        for n, d in (("A1", self.d), ("B1", other)): Employee.objects.create(nik=n, name=n, gender="L", join_date=date(2024, 1, 1), department=d)
        self.login("adm"); body = self.client.get("/api/employees/export/").content.decode("utf-8-sig")
        self.assertIn("A1", body); self.assertNotIn("B1,", body)

    def test_schedule_leave_requests_exports(self):
        Employee.objects.create(nik="S1", name="Sat", gender="L", join_date=date(2024, 1, 1), department=self.d, status="aktif"); self.login("hrd")
        for url in ("/schedule/?export=1&format=xlsx", "/leave/?export=1", "/requests/?export=1&format=xlsx"):
            r = self.client.get(url); self.assertEqual(r.status_code, 200, url); self.assertIn("attachment", r["Content-Disposition"])
        self.assertIn("S1", self.client.get("/schedule/?export=1").content.decode("utf-8-sig"))


class PoliImportExportTests(BulkBase):
    def test_medicine_and_diagnosis_import_export(self):
        self.login("poli")
        r = self.post("/poli/medicines/import/", [["code", "name", "unit", "min_stock"], ["PCT", "Paracetamol", "tablet", "50"]]); self.assertContains(r, "1 ditambahkan")
        m = Medicine.objects.get(code="PCT"); self.assertEqual((m.stock, m.min_stock), (0, 50))  # stok tidak diimpor
        self.post("/poli/medicines/import/", [["code", "name", "unit", "min_stock", "stock"], ["PCT", "Paracetamol 500", "tablet", "20", "9999"]])
        m.refresh_from_db(); self.assertEqual((m.name, m.min_stock, m.stock), ("Paracetamol 500", 20, 0))
        self.post("/poli/diagnoses/import/", [["code", "name", "category"], ["J00", "Common cold", "Pernapasan"]]); self.assertTrue(Diagnosis.objects.filter(code="J00").exists())
        self.assertIn("PCT", self.client.get("/poli/medicines/?export=1&format=csv").content.decode("utf-8-sig"))
        self.assertEqual(load_workbook(io.BytesIO(self.client.get("/poli/diagnoses/?export=1&format=xlsx").content)).active["A2"].value, "J00")
        self.assertEqual(self.client.get("/poli/reports/stock/?format=xlsx").status_code, 200)

    def test_poli_import_forbidden_for_others(self):
        for u in ("hrd", "adm"):
            self.login(u); self.assertEqual(self.client.get("/poli/medicines/import/").status_code, 403, u)


class AuditExportTests(BulkBase):
    def test_audit_export_superadmin_only(self):
        self.login("sup"); r = self.client.get("/audit/?export=1&format=xlsx"); self.assertEqual(r.status_code, 200)
        self.assertEqual([c.value for c in load_workbook(io.BytesIO(r.content)).active[1]][:3], ["waktu", "pengguna", "modul"])
        self.login("hrd"); self.assertIn(self.client.get("/audit/?export=1").status_code, (302, 403))
